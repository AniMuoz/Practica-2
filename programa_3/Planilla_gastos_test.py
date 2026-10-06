import base64
import datetime
import io
import os
import sys
import threading
import webbrowser

import openpyxl
from flask import Flask, flash, redirect, render_template_string, request, send_file, send_from_directory, url_for
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

app = Flask(__name__)
app.secret_key = "planilla-gastos-test"

data = [[], [], [], [], []]

fecha = datetime.date.today()
dia = str(fecha.year) + str(fecha.month).zfill(2) + str(fecha.day).zfill(2)
fecha_pro = f"{fecha.day}/{fecha.month}/{fecha.year}"

topicos = [
    "CONSUMOS BASICOS",
    "TELEFONO E INTERNET",
    "GASTOS COMUNES",
    "ARRIENDO DE OFICINA",
    "COMBUSTIBLE",
    "ESCRITORIO Y OFICINA",
    "ESTACIONAMIENTO",
    "ARTICULOS DE ASEO",
    "GASTOS DE REPRESENTACIÓN",
    "VESTUARIO Y CALZADO",
    "PASAJES, PEAJES Y CORREOS",
    "MANTENIMIENTO, REPARACIÓN Y SEGURIDAD",
    "EQUIPAMIENTO",
    "ALIMENTACIÓN",
    "OTROS",
]

meses = [
    "ENERO",
    "FEBRERO",
    "MARZO",
    "ABRIL",
    "MAYO",
    "JUNIO",
    "JULIO",
    "AGOSTO",
    "SEPTIEMBRE",
    "OCTUBRE",
    "NOVIEMBRE",
    "DICIEMBRE",
]

PAGE = """
<!doctype html>
<html lang="es">
<head>
  <meta charset="utf-8">
  <title>Gestión de Gastos</title>
  <link rel="icon" type="image/x-icon" href="/favicon.ico?v=4">
  <style>
    body { font-family: Segoe UI, sans-serif; margin: 24px; color: #1a1a1a; position: relative; }
    .logo { position: absolute; top: 0; right: 0; height: 72px; width: auto; }
    h1 { margin-bottom: 4px; }
    .muted { color: #555; margin-top: 0; }
    form.grid { display: grid; grid-template-columns: 160px 280px; gap: 8px 12px; max-width: 480px; align-items: center; }
    label { font-weight: 600; }
    input, select, button { padding: 6px 8px; font-size: 14px; }
    button, .btn { background: #1f4e79; color: white; border: 0; border-radius: 4px; cursor: pointer; text-decoration: none; display: inline-block; }
    button.quitar { background: #8b2e2e; }
    td.fila-acciones { width: 1%; white-space: nowrap; padding: 4px 6px; vertical-align: middle; }
    .fila-botones { display: inline-flex; gap: 6px; align-items: center; }
    .fila-botones form { display: inline-flex; margin: 0; }
    .fila-botones .btn, .fila-botones button { padding: 4px 10px; font-size: 13px; border-radius: 6px; font-weight: 600; }
    .fila-botones .modificar { background: #e7eef5; color: #1f4e79; }
    .fila-botones .guardar { background: #1f4e79; color: white; }
    .fila-botones .cancelar { background: white; color: #4d5966; border: 1px solid #c5ced6; }
    .fila-botones .quitar { background: #f8e8e8; color: #8b2e2e; }
    tr.editando { background: #f4f8fb; }
    td input, td select { width: 100%; box-sizing: border-box; }
    .acciones { display: grid; grid-template-columns: repeat(auto-fit, minmax(240px, 1fr)); gap: 12px; margin: 16px 0; }
    .accion { display: flex; flex-direction: column; gap: 6px; padding: 10px 12px; border: 1px solid #d5dde6; border-radius: 8px; background: #f7fafc; }
    .accion h2 { margin: 0; font-size: 15px; }
    .accion p { margin: 0; color: #4d5966; font-size: 12px; line-height: 1.3; }
    .accion input[type="file"], .accion select { width: 100%; background: white; padding: 4px 6px; }
    .accion .btn, .accion button { margin-top: auto; text-align: center; padding: 6px 10px; border-radius: 6px; font-weight: 600; }
    .accion.guardar .btn { background: #1f4e79; }
    .accion.recuperar button { background: #2f6f4e; }
    .accion.procesar button { background: #8a5a12; }
    table { border-collapse: collapse; width: 100%; margin-top: 12px; }
    th, td { border: 1px solid #ccc; padding: 6px 8px; text-align: left; }
    th { background: #e8eef4; }
    .flash { background: #fff4d6; border: 1px solid #e0c36a; padding: 8px 12px; margin: 12px 0; }
    footer { margin-top: 28px; color: #666; font-size: 13px; display: flex; justify-content: space-between; align-items: center; gap: 12px; }
    button.apagar { background: #3d4652; padding: 6px 12px; border-radius: 6px; font-weight: 600; }
  </style>
</head>
<body>
  <img class="logo" src="{{ url_for('logo') }}" alt="Logo">
  <h1>Gestión de Gastos</h1>
  <p class="muted">Berfre Ltda. · Hoy es {{ fecha_pro }}</p>
  <p class="muted">Codigo de dia: {{ dia }}</p>
  {% with messages = get_flashed_messages() %}
    {% if messages %}
      {% for message in messages %}
        <div class="flash">{{ message }}</div>
      {% endfor %}
    {% endif %}
  {% endwith %}

  <form class="grid" method="post" action="{{ url_for('agregar') }}">
    <label for="topico">Tópico</label>
    <select id="topico" name="topico" required>
      <option value="">Seleccione</option>
      {% for topico in topicos %}
        <option value="{{ topico }}">{{ topico }}</option>
      {% endfor %}
    </select>
    <label for="proveedor">Proveedor</label>
    <input id="proveedor" name="proveedor" required>
    <label for="nboleta">N° Boleta</label>
    <input id="nboleta" name="nboleta" required>
    <label for="fecha_boleta">Fecha boleta</label>
    <input id="fecha_boleta" name="fecha_boleta" type="date" required>
    <label for="monto">Monto</label>
    <input id="monto" name="monto" inputmode="numeric" required>
    <span></span>
    <button type="submit">Agregar</button>
  </form>

  <div class="acciones">
    <section class="accion guardar">
      <h2>Guardar</h2>
      <p>Descarga las boletas cargadas, sin armar el detalle del mes.</p>
      <a class="btn" href="{{ url_for('guardar') }}">Guardar datos</a>
    </section>
    <form class="accion recuperar" method="post" action="{{ url_for('recuperar') }}" enctype="multipart/form-data">
      <h2>Recuperar</h2>
      <p>Carga un Excel guardado antes y suma esas boletas a la tabla.</p>
      <input type="file" name="archivo" accept=".xlsx" required>
      <button type="submit">Recuperar datos</button>
    </form>
    <form class="accion procesar" method="post" action="{{ url_for('procesar') }}">
      <h2>Procesar</h2>
      <p>Elige el mes y descarga el detalle de gastos listo para revisar.</p>
      <select name="mes" required>
        <option value="">Seleccione el mes</option>
        {% for mes in meses %}
          <option value="{{ mes }}">{{ mes }}</option>
        {% endfor %}
      </select>
      <button type="submit">Procesar datos</button>
    </form>
  </div>

  <table>
    <thead>
      <tr>
        <th>Tópico</th>
        <th>Proveedor</th>
        <th>N° Boleta</th>
        <th>Fecha boleta</th>
        <th>Monto</th>
        <th class="fila-acciones"></th>
      </tr>
    </thead>
    <tbody>
      {% for fila in filas %}
        <tr class="{% if editar == loop.index0 %}editando{% endif %}">
          {% if editar == loop.index0 %}
            <td>
              <select name="topico" form="editar-{{ loop.index0 }}" required>
                {% for topico in topicos %}
                  <option value="{{ topico }}" {% if topico == fila[0] %}selected{% endif %}>{{ topico }}</option>
                {% endfor %}
                {% if fila[0] not in topicos %}
                  <option value="{{ fila[0] }}" selected>{{ fila[0] }}</option>
                {% endif %}
              </select>
            </td>
            <td><input name="proveedor" form="editar-{{ loop.index0 }}" value="{{ fila[1] }}" required></td>
            <td><input name="nboleta" form="editar-{{ loop.index0 }}" value="{{ fila[2] }}" required></td>
            <td><input name="fecha_boleta" form="editar-{{ loop.index0 }}" type="date" value="{{ fila[3] }}" required></td>
            <td><input name="monto" form="editar-{{ loop.index0 }}" value="{{ fila[4] }}" inputmode="numeric" required></td>
            <td class="fila-acciones">
              <div class="fila-botones">
                <form id="editar-{{ loop.index0 }}" method="post" action="{{ url_for('modificar', indice=loop.index0) }}"></form>
                <button class="guardar" type="submit" form="editar-{{ loop.index0 }}">Guardar</button>
                <a class="btn cancelar" href="{{ url_for('index') }}">Cancelar</a>
              </div>
            </td>
          {% else %}
            <td>{{ fila[0] }}</td>
            <td>{{ fila[1] }}</td>
            <td>{{ fila[2] }}</td>
            <td>{{ fila[3] }}</td>
            <td>{{ fila[4] }}</td>
            <td class="fila-acciones">
              <div class="fila-botones">
                <a class="btn modificar" href="{{ url_for('index', editar=loop.index0) }}">Modificar</a>
                <form method="post" action="{{ url_for('eliminar', indice=loop.index0) }}">
                  <button class="quitar" type="submit">Quitar</button>
                </form>
              </div>
            </td>
          {% endif %}
        </tr>
      {% else %}
        <tr><td colspan="6">Sin boletas cargadas.</td></tr>
      {% endfor %}
    </tbody>
  </table>
  <footer>
    <span>© 2026 Berfre Ltda. Todos los derechos reservados.</span>
    <form method="post" action="{{ url_for('apagar') }}">
      <button class="apagar" type="submit">Apagar programa</button>
    </form>
  </footer>
</body>
</html>
"""


def fecha_iso(valor):
    if isinstance(valor, datetime.datetime):
        return valor.date().isoformat()
    if isinstance(valor, datetime.date):
        return valor.isoformat()
    texto = str(valor).strip()[:10]
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(texto, formato).date().isoformat()
        except ValueError:
            continue
    return ""


def filas_vista():
    return [
        (data[0][i], data[1][i], data[2][i], fecha_iso(data[3][i]), data[4][i])
        for i in range(len(data[0]))
    ]


def leer_boleta():
    topico = request.form.get("topico", "").strip()
    proveedor = request.form.get("proveedor", "").strip()
    nboleta = request.form.get("nboleta", "").strip()
    fecha_boleta = request.form.get("fecha_boleta", "").strip()
    monto = request.form.get("monto", "").strip()
    if not (topico and proveedor and nboleta and fecha_boleta and monto):
        return None, "Debe completar todos los campos."
    try:
        fecha_boleta = datetime.datetime.strptime(fecha_boleta, "%Y-%m-%d").date()
    except ValueError:
        return None, "La fecha de la boleta no es válida."
    try:
        monto = int(monto)
    except ValueError:
        return None, "El monto debe ser un número entero."
    return (topico, proveedor, nboleta, fecha_boleta, monto), None


def workbook_sin_procesar():
    guardar = openpyxl.Workbook()
    hoja1 = guardar.active
    for i in range(len(data[1])):
        hoja1.cell(row=i + 1, column=1, value=data[0][i])
        hoja1.cell(row=i + 1, column=2, value=data[1][i])
        hoja1.cell(row=i + 1, column=3, value=data[2][i])
        hoja1.cell(row=i + 1, column=4, value=data[3][i])
        hoja1.cell(row=i + 1, column=5, value=data[4][i])
    return guardar


FORMATO_PESOS = '_-[$$-340A]* #,##0_-;-[$$-340A]* #,##0_-;_-[$$-340A]* "-"??_-;_-@_-'
FUENTE = "Calibri"
RELLENO_TITULO = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
BORDE_FINO = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def borde_fila(hoja, fila, columnas):
    for columna in columnas:
        hoja.cell(row=fila, column=columna).border = BORDE_FINO


def combinar_etiqueta(hoja, fila, texto, negrita, centrado, relleno=False):
    hoja.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=4)
    celda = hoja.cell(row=fila, column=2, value=texto)
    celda.font = Font(name=FUENTE, bold=negrita, size=12)
    celda.alignment = Alignment(horizontal="center" if centrado else "left")
    if relleno:
        celda.fill = RELLENO_TITULO
    borde_fila(hoja, fila, range(2, 5))


def celda_monto(hoja, fila, valor, negrita=False, relleno=False):
    celda = hoja.cell(row=fila, column=5, value=valor)
    celda.font = Font(name=FUENTE, bold=negrita, size=12)
    celda.number_format = FORMATO_PESOS
    celda.alignment = Alignment(horizontal="right")
    celda.border = BORDE_FINO
    if relleno:
        celda.fill = RELLENO_TITULO
    return celda


def workbook_procesado(mes):
    guardias = openpyxl.Workbook()
    hoja = guardias.active
    mes_txt = mes.upper()
    hoja.title = f"GASTOS {mes_txt} {fecha.year}"[:31]

    hoja.column_dimensions["B"].width = 59.43
    hoja.column_dimensions["C"].width = 13
    hoja.column_dimensions["D"].width = 11.29
    hoja.column_dimensions["E"].width = 19.14

    hoja["B3"] = "EMPRESA: BERFRE LTDA."
    hoja["B3"].font = Font(name=FUENTE, size=11)
    hoja["B4"] = "DIRECCIÓN: 22 NORTE 1150 SANTA INES, VIÑA DEL MAR"
    hoja["B4"].font = Font(name=FUENTE, size=11)
    hoja["B6"] = f"DETALLE GENERAL DE GASTOS MES {mes_txt} DEL AÑO {fecha.year}"
    hoja["B6"].font = Font(name=FUENTE, bold=True, size=12)

    combinar_etiqueta(hoja, 8, "CLASIFICACION DEL GASTO", True, True, True)
    monto_titulo = hoja.cell(row=8, column=5, value="MONTO")
    monto_titulo.font = Font(name=FUENTE, bold=True, size=12)
    monto_titulo.alignment = Alignment(horizontal="center")
    monto_titulo.fill = RELLENO_TITULO
    monto_titulo.border = BORDE_FINO

    for indice, topico in enumerate(topicos):
        combinar_etiqueta(hoja, 9 + indice, topico, False, False)

    fila_total = 9 + len(topicos)
    combinar_etiqueta(hoja, fila_total, "TOTAL GASTOS DEL MES", True, True, True)
    celda_monto(hoja, fila_total, f"=SUM(E9:E{fila_total - 1})", negrita=True, relleno=True)

    firma = fila_total + 3
    hoja.cell(row=firma, column=4, value="CRISTIAN BERNAL P.").font = Font(name=FUENTE, bold=True, size=12)
    hoja.cell(row=firma, column=4).alignment = Alignment(horizontal="center")
    hoja.cell(row=firma + 1, column=4, value="JEFE DE OPERACIONES").font = Font(name=FUENTE, bold=True, size=12)
    hoja.cell(row=firma + 1, column=4).alignment = Alignment(horizontal="center")

    totales = {}
    fila = firma + 4
    for topico in topicos:
        filas = [j for j in range(len(data[1])) if data[0][j] == topico]
        if not filas:
            continue

        titulo = hoja.cell(
            row=fila,
            column=2,
            value=f"DETALLE GASTOS EN {topico} MES {mes_txt} DEL AÑO {fecha.year}",
        )
        titulo.font = Font(name=FUENTE, bold=True, size=12)

        encabezado = fila + 2
        for columna, texto in enumerate(("PROVEEDOR", "N° BOLETA", "FECHA", "MONTO"), start=2):
            celda = hoja.cell(row=encabezado, column=columna, value=texto)
            celda.font = Font(name=FUENTE, bold=True, size=12)
            celda.alignment = Alignment(horizontal="center")
            celda.fill = RELLENO_TITULO
            celda.border = BORDE_FINO

        primera = encabezado + 1
        for desplazamiento, j in enumerate(filas):
            actual = primera + desplazamiento
            proveedor = hoja.cell(row=actual, column=2, value=data[1][j])
            proveedor.font = Font(name=FUENTE, size=12)
            proveedor.border = BORDE_FINO
            boleta = hoja.cell(row=actual, column=3, value=data[2][j])
            boleta.font = Font(name=FUENTE, size=12)
            boleta.alignment = Alignment(horizontal="right")
            boleta.border = BORDE_FINO
            fecha_boleta = hoja.cell(row=actual, column=4, value=data[3][j])
            fecha_boleta.font = Font(name=FUENTE, size=12)
            fecha_boleta.alignment = Alignment(horizontal="center")
            fecha_boleta.number_format = "DD/MM/YYYY"
            fecha_boleta.border = BORDE_FINO
            celda_monto(hoja, actual, data[4][j])

        ultima = primera + len(filas) - 1
        total_fila = ultima + 1
        combinar_etiqueta(hoja, total_fila, "VALOR TOTAL", True, True, True)
        celda_monto(hoja, total_fila, f"=SUM(E{primera}:E{ultima})", negrita=True, relleno=True)
        totales[topico] = f"E{total_fila}"
        fila = total_fila + 3

    for indice, topico in enumerate(topicos):
        referencia = totales.get(topico, 0)
        celda_monto(hoja, 9 + indice, f"={referencia}" if referencia else 0, negrita=True)

    return guardias


def excel_descarga(libro, nombre):
    buffer = io.BytesIO()
    libro.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


FAVICON_ICO = """AAABAAMAEBAAAAAAIACwAwAANgAAACAgAAAAACAAFAoAAOYDAAAwMAAAAAAgAK8RAAD6DQAAiVBORw0KGgoAAAANSUhEUgAAABAAAAAQCAYAAAAf8/9hAAADd0lEQVR4nHWTf0zUdRjH35/P9/u9S+8OL+XHeVzGdZrDu4DGKoZw2jTIQMniKghb05NOMdcvldPasaZYjmqupKYi0RY1HOEpbWKLa5aB9Qe7GQULg2CBoAhHynHe9/t9+oOD9Y/PX897e573s+d59gJmg7tcLkEUeExiJ4Ctc4JzvArAAwBgQHl5pgRgvlhgjM3l6x960BKo/3A7vVNRGE1YvKDYaJC2ADi4x533c+H6jO8BJAFAbJiAWHOuJcnYfOitZ+nf7s+IqDUa7TtJqSvuUwC0jfz6aRVRW5iGGqi9yXu9rCjrIIAMxgAwhuNP5tovTFypJYr4ZRo4RUertpDFvPT03oLE7d96rV2eguV0yPuiGvr9hELhFrrR9TFVejbc0S/UNoqcM9b953CL0eG43Nrgf7vyfX87xvqah46tSoEoHQZjS9atmFHq2s8LL3s6WWFRHm315Mn25WZJVSk0t/wHAJw6oO2W3z4N4R43pmUrTSuQFSLOGBOMGiA8g731fytnr+qF3v7RagAHmE6S7Lej0dKKEqd1Zjq0Zm3iiLk0Jx5cp1HkKZkDBNEgYTo0w764OI7gZLzqyM5i/tbOS9/99MdLYEDXm+78WqIfBuT+z6naWxZ15VrUxlcSiBodKjWlU1OFiYpzktXD+8voVm8dEZ1XO057Sb9Q08dEUXxkpTXxQOnT2UXlJWsQn3Y/RoJDOP5lgEZ+62BTU1OjyzIex77XNiXda0ug0eAgjp66gKZzHc3DNyZr5v6fAuCbAufK7pr9L5Ay2KAMdnykrs60TTBROhtorCyh8a97jnhLFffGVNVpN8oAEjmfPSE3GMSsBVpxJ/XsqT9ZtY6K89OiJXl26jyxebLgseSndBLOrE03Xandna5Gfnxefc+dFtFqtTYiHxcBqEyGGo7IVizWabbtyFXy065CH69X+/tvGn7pGRseb3lmSYTBEZeouwMN14zdDGsikQgk8V2VA2DxSw29jEHdtaP5iWvBfwRLqokZzYvknoEJfj0UtWlTFt2OWxanDP81qdlXfZkaA0N1SUm4FpVpngHEcMi2meNaa3bnkBx8gwKfbFYArKb25y4eez2THjDpzgF4FHcJLvB5ww3Oh5Mvbdu4ihjwldWkOwNg0yxEDC5AuKuJz+fjgsBi4GIXAC8ASeAMPh/4/zEGgP8AAplsLvgmI0oAAAAASUVORK5CYIKJUE5HDQoaCgAAAA1JSERSAAAAIAAAACAIBgAAAHN6evQAAAnbSURBVHicpZd7cFXFHcc/u+fc901CwlO0IIrIGBRLhVotIqgNIoIR03asgFi1QHkoylsaiYKaWoGOAYoKaIxphaKIOlSFKq0IpRR5FlDBCimREExCcu/NvWf31z9yAyE6ZWy/M2fmzJ6z+/3ub/f3gnPDAVBK0a1bh45to9GeLcdbvfcA2imlWo//T3DSvBQWFviBccC/gNpAwLk+/Y+bfgCuAzygFhhfWFjgdxxNQUGBA+hvQ6yblTuOBrgD2NkuOyqL5t4lU8bcJEAsKxIY2GLOdUohSqmFwIQ2GSEBdgI3AmitEEEBiv8CdZpYa1zXHQBsCAV88uDPfyTHti305MtXrFSWmYkjBwnQEAq53wf6KRBHqcUAknzjkaq9SxonjBwktw+56pMNpVMvBbq34HFaC1Gkzai1IhKJ5ALlgNw57GrZ9958I8fLjRxeLok9SySxb6lI7Wozc+wQC9T6XB0PBZxirRX1u0qekcoykc9XGKko9eTkH2Rj6cPJnhefd/jjPxZ26NQp2j5tVdKcygWEprPrZK1Ma2hoGHfjtZcFfzV5uPQf0MvSmHJSJ2oREQLRIAQDdseH+/j86Anlc1RmypNEyrM5bz03sSRy6fnjOVbjJesTrnI0+lRc+ve9xLfrvfkXbtmyf3dlZb0DlF1wQU5xxdGTFZLefVRrPQmk6OorL3Ie+HmeFAy/2oJyvFMxrLH4Q37ICNsvPq2QJ0vecla8uolEyqwFHno0v8PVWw/FX/7gE48Rg68y8x4e7nznks5QFyfVmMJxtGitEVer1/70Dx5btJaP/3nkBPA0UKKACq1VB7/j3hvft3QwXdr91DtU6VljXZ/fRWVFpK6q1ix4fr377Mp3OVET2wJMl9rb9rLh8BQ8+QWQ9cZHXzHt5aPuoVo/D9+Xx8yxeWR0zEZqGvA80yQkM0yqPm5feX2L89D8P3CypqFCAcVKMVlrZ7BW/Ll6+6J1Gee1HYoxxsaTsuyV991f/+4tDh2p/gyYKX8dtp76iru8441FbshpR8xgjOBEXUjBcxurmF3+b7xAG4oevJXxPxuAjgQxNTEaG1OEz28rG9f/3eSPLXEb4olfN1+IR5RSAtzhaI38++VX3l45RfrkdhGgGpg0uF/3THm/b74p7XVAVvUWKb1czIrclFnZy9qVvSS1PFfsyl4iq3tLw4rL5dH8dhJQyEXdzpdViyeI/GulSM0qefXZcTZ972YqpU57gAc8qpSaIyKTgRvSPlwCPC27+/dgV00xPn0tKYtNWg+Uk55/FowRXFdB1OVYZZJHyo9Q+tdTDL6hNzf1z7UPFJVba+UxpVSRiLjNwcEHJIG1jlbDjBXaZ4dvPn5y9EFWby6m0Y7AUdiEMSil1DkimwgkPSEY1tAhyDvvVpNXfMgAjlK8CepWEfEDKZ22QBKYpLUaaqzMjgR9f6v6KrZixi2/P4DHCALa2oSxVnDORW6soBQEs33UxYUZCz/jjgWHGNCvh/OrScOsCDcoJROVUknA1UAKmKC1WmStjBKRxfUHn8t48el7Oy3fpt02o3aakjePa+VztRt1sQLWfp3YWrACbtQF16Fk3Zd0++UeXtgKTz16NxvKpjG3aKRaNn90yFp+KyLjlFIpBUzXSs21IgUisjl+YNmuUGa4M1qbhuo6p/i5d3hq6XouiMQpHtmF23+QA2LxYobmrCciuGEHlGbNRyeZVvoFRxtCTB87mGn3/4hI2yxsbQONiRShTtmy5tW/2NHTXnASydQMrZSaJaD6Xt61LUdf2ho6L6dz6qt6k6yucyIBH3NnF3D4g3kMzBvEiIVHuWbGHjb/M4ab6cPxKxy/ws10zYf76s01M/YwYuFRBuYN4vAH85g7u4CI30eyqhZrLKG2GRBvtD6/64SDfoyR6QroqJSa6vc5D13Xtwczxw81Awf1VqQ8naqPIyL4w0GIBNj78WdMe+p13t6wnWG9g/xmVFcQmPvacVZtq2fgtb3s0zNH6NwrL4KGRpKxBEopfNEQ+B02btwlRYvWqg+2HogBzwILWrpRd+Bx4CcFQ/oyZ/IwLu/THeoTZy/kc9j4/m4KF6xj564DKKWkLu6tWfXbsZl33Jd3E8drvVRDwm0pfPeOz5jx1Br79sYdClidkZExp77+1AGRJuJmNyQnGr1MKZUAPgJkTMF1tmLrMyJV5SKfvtCUDfcuFYm9JvOm3CZA3O93JezXs3KywsT2Ll0vx8pEPnneSFW5VGx9Ru75yfUCSnK7RGw46EqLispHi7Ss088PtFbe0KHfC88cddXjvS/OFp/Pl5ozabic2rtEpLJM5Hi5zBp3iwVk0+IRR0YOvrQAEAVFAb+L1K0uj336vJ09cbjx+wNycaegrHuiv1T/cbgNB1wBvl9YeJqPlgIA+rmOFiAs9Y/fIwenS8nU61M5GT7JaZMhLy64X+ZOHi6AvLvkdrOxJD+VFQncEAg4N5EOr+1zMu5um5NlssPKLJ7cR+x7PxbZ+jP5vHSIDfodAa5KO4+GM/Vcayj2VwUI+hg/+nuMuvUynly5jTEPLsMCa4pv4cabe1K6/G/Wcenqxe3yeWNyy578/YH5sViMB0Zcwow7bySaHcDWNWLrhG+I2nyTAAEsYLFKi4BXEyPqd3l86kDuG3oZ1XUJ+lzZWczxenVF93Zet46ZB7fXVjmzftHbG359FxsB78Kebf3UNWJqGk+fsm6ycXMi+prpm3Gmgu3dsVZpUCjECOZEPV27tKHPFZ1prEkoJyNgnnxpe3j7warJgGmobHByc9vpCy/I0KY6jhgBlS40s/xiLUZEtNtq080CmqujCoEjwJ7CsWt+jBXjZgS0iKCUwjZ62FgSx1EgQijoA4gDuD4tNu5hkgalmqKjE/Wjg655+bVPVd6sTW7Ss0fcIMfSFbK0FgBwxBjbRyn1RtGKbUOuvbPMeecvh7XODBgd9iFWzqpnpcmRT48oBWJBh1x0ht/8aXMFP5ywwRn5xNZT+4+cekKEPokEX7TkbHkEzVY4ISIPBgJcsnl35fK8ia/b/IlrnR37vrROdshqv4tnbKuTRCVTVim/g5sdsDsOnLT50zc5g6dvsh/uPbG8fWbgu1qpWcCJNMfZs1vhrDI95Lr9gHWuo+T+/F7y+bp7PNk1xcj+qalRN/cUoAxAtt217Nib+ebeId3E1UqAdSHX7af1aQO5nKMxaY0zHVLTIkOALVlRv8y+p6+kdk6Jj8vvZYElowd0Dd52TeePOmQFhKYoOsQ5Q/ytW7NvEqIVIFKogTHAJz2+00Zyu+UIsF8rVUlT33i3iDS3pl+Ldv8v0l0yDBjQPgpMBSqBGmDOFVd0jJxpiv+/rvhcOH0/ckKh84HOrb99G/wHY6iJK84cKnUAAAAASUVORK5CYIKJUE5HDQoaCgAAAA1JSERSAAAAMAAAADAIBgAAAFcC+YcAABF2SURBVHiczZp5mFTF1cZ/VXVv9/TsDJssKoKAiCZCBCEBNCIqIKhRRNSIigJRQPz4FBWChsUFFRVRUVxiDIuoUUHAqAgoigtGRJBFkAxgWIZl9pnue6vO90d3wzAiwazfeZ5+nu7bXVXve06dU29VNfzjpgDvMM/1Edoc7jsv1dd/1Ez6TSwWa+L75jGj9TKgVe3vD9OmFbAEeARocrg+/51Wc5DcSMQbA+wFBBCt1TbgpNT3NaOTft8aKEz/3hhdGol4Ywsg99+MG83B8Juo7w9SSm0CJBb1ZezwPuHQAWcFgHhGb00BhSThNOnWwJYU+MeUUuM5SH5TNOpf/+jw86Nz5/ZLtznSVPxRwA94PRIxfbVWn6YGlasu6hSsfWeSk12zRL57UYYOODMExPPMVg5GIg3+W0CUUk9onZzyH79y50s3XN4tSPcH3PWvAA3JpDoAPMPzOnueXpgeqEeXtuHyl++wsuOPIoW/l8Ta6RJsnJEi0S0ERCu1FWiWehWmwD9uTNKxlWuefEL2zhHZMdN+NPeO4PLeHcInJv56jcjSU3NzYx3r5WW2l7lzDeBzhESv/TANPASIRCKtcfaORGivBlS7k49z427uy0U9O2i0IiyrQimF1goRwSmFF/UZNuYF++SsZQZki3PigBZKMU2E4QB2/dOP6YKcYeG+UotgvJwYKISor4I9pdUF7UdEyyvjXwE/rYXPAC7lyO8RMIAFyIKGCd+MCgL7GyD7+CZ1uePGC+x1/boYPyeGLalEREh7E8Bah1IKU5BN9b4yTu/zu+DrTX/zlVIgsvqm807sOGNJYbxs1bRpXn7WTa6sKlSp5LbOpSE5LyOiZ7++QrKyohV9L+16X1bjq7QVyfW8yAMVFRW7a2NNezzNKuZpPSR07jagUX5uJiOv62FHXNPD1GlUgJRUYK07BLhzLtljbiZSHfDiax/ZSdPm602FuxVCtVIo64iCrB4z5PwNEydc3Y/QWltSaVCg9aG5KiKis2Okwkq77newat02lFI7fd88kEiETwJVKey6ZgQGKKV+KyJtjFYMvuKX4W1De5pmLZsoyqsI4gHGMwdC5pwgInjZMVDw1uJVbvyjb8iKv2w2AJ5Wfwqd3A5UNqvvzy7cE3YVETq3a+7uGnmRPq/7aSAQlh+chjWj6UTEi/puR1EJk6a9Kc/MWeYFoUUptU4pNck5NzMdgbOBu7RW3USErqe3Kpx+37XHt2lzrCMR6KAyjjEapVTaQ1jr8DOjkBFh5ScbZcJjb7h5764yANGIXm6dHhuG4TKRHlm8uvsWYMgHq8oajXtllyxdX+EBXHxuO8YM68vPOraC6gS1x0mPpSM+xCJ8sfIbmfjYfPenP3+eLi7vA79T6YRQSq0xxgwKgr2b2Tn/dZx0CUsrQ2O0l+wMrLX4UR+yY2xev03un77APv/Sci90Ds/odVHfjK+oDuZkeIriGSeP9Iy61cRMY6osRDRY4eVPSpnw2m6+2pacBUOv6MatQ3vTvHVTDkTaGNI8RATnBC87AwQ+/HSDG3z787Ju8w6jlEKTqr0ikgc4pfL2vj7nnT4unvjAK8j2RCS01qG1wq+by97SSsZNnG3b975bzZj9vufE7ciKebeE1p1cUR3MCV9o27/s2bZro9new0arxq48DJ1FwkqLSwj9uuTz6fgWTBvYhCZ1PKbPWsZPzx/H3ZNms6+0Er9uLlorrE3mllIKYzTxkkrwDVHfU8Wl6RRQd6fjdRtwP1AejZq+8bhdsuW1u/Kbtm26wMvL/Dmhc/HyKv3MSx/Y+59402zbuR+gPCfmP1FWFYwHKmTmKWeGARO8mO6KBRe3FpRW6tBSbZ3gGQWZhj17Aqa+vZcpC3ZTkRCaHlOHO2/qzXWXdSOam4krrTwQAb9uLksXr5Je1z5sq+KBp7W+3Tl3f1pRhsDtwL1ASTRqLorH7VIRMeycuWDhwpXn3D75Fflqw3YPcFlR78WKeDgO2Cp/atcmURJMjETUrzAKV2UtSil1BCkggLWCH9GQodm8rZrJb+7l6ff2AHBam2MZO7wvl/TuAEpBxGPpki+lz6BHbXll3PM8PSYM3T2Apzh08RoN3AdUZHhe7+owrNBKPehEzgSIRb1FVXH1Wwg+l41d6ic+Kr7baHW9ydARV2EFxCmljlpViiSrmZehwdd8tqGCSW8U8cbKEgDO6nQS946+lHh1IBcMekSVV8bxPH1nGLp7U4636fDWJHErMBEIlVIREfGUUu+J6AfBLhKZa+LPjR/j+epmE9MFVFqcE/tjgNc2J8lk9WIGRFjwl3IefWsP76wpwzca42lXHQ8TntZ3hc5NToMH5HArcQawBmihlMIzelfUU5dGPVm+5/enDK4qDe+IZXvNiDts4KzSSqt/wYZEBAIrRDM05Hks/6SUAU9uZ/veuAOUMbrQWtcGqK6B9YBO16kH+cCCFPilIqKC0HYLQpZ2aptTuH1HonnTY6NQGtpEINozyvyzyEWSie1HNNFsw7otVUx68jtmfrgPgK4dWmqFcu9/trGZb/TSwLrzgeIUZpeeOmnw84CuWqt3nZPesnWueWjGa2unv7D4hE1bd2MUdmSvBmpUz7q6UcMIVFpCKxj9j9GoWZF27Erw8Ft7eXTRbhIWmh9bn9tv7M2g/t1IhJYLBk6xi1esM76vVwSB65UiccCBeSQ9/wut1TvOyUXy0UNi6+a8YY6p06N062477YXF6sEZf9b7Syqon2MY3bchQ8+uQ1aOh0vmwVETsU7QWqEzDRVlIdPf28/983ZRVGapk5fFrYPP48aru5PXIB+7vwId8ahKhFx47ZTw3RXrPaP1CutcT6BE1fS8UmqxiPSVLc+7IAjn+/lZ5ySKSl0k6mtyYmzbvIOHnl7EtBeXYK2ldaMoYy9qyJU/z0NFFLYiufjoHyigKd2HydJIQpj5UQkTX9/Fhh1xjDEM+/UvGTW4J8e2aARlVQSJAM8zOOcwEZ/yeMCF10wJ3/tkg6e1Wu6c9FHAauBUpdRnCjpa5yJV655aECvIOceVVIRKKe+A/olFIRbhq1Xfcs/j85gz/zMAurTK4rcXN+TcdtlJgVZlkwItFZBDqoyCt78oZ8Jru1i+sQKAy/t04M6b+nLqac2hKkFQ9X39FQSWaEEOBCHte44LVq3b6iul1irgS+AnWqtq52TG6jfvPvHUHu16sn1PGCZCT2v9fV2SlQFGs/T9r5g4dR6LV6wH4OIO+Yy5sD4/a5UJCUcQT7rcj2qIaD7fUMmkeUW89lkxAN07n8TYEX05q9upYB1hRTVaqxrAU/orIwLZGWz86q8y+elF9tWFn3sl5VWArFYkt2zXKxiLUo3rF2Rzbb8uwagbevr1m9aD0krC0NbaAwiIYHIzIQh5ecGnTJw6n9UbtgMwpHs9brugLs2bZgCweVu13D+vyM1Yuk8B+ietmzJ2RB/69e4IvoctrYTDSGrPM5CbSdG2IqY8+2c77YXFprwyDrCV5IL7TM2sKwCGATcDBY0b5MutQ3oy5MqziOVnK1dSkUzUWrswrRU6N5N4aSXPzFnGvU8s5Ltd+8n0YXSfhgjCw4v2UlJlaXVCQ0Zcc44M6t+NjLxs5UqP0GdeFlXF5Tw9aymTpy+Sv+0uVkCx75upQWAfBfZBStJxcB0gGo2eIDa8JxHa/oBq26oJY4f3cZf3OUMflbe2FjF91hIefu5d9peUC+BAFUc8NfnrtybEW5xx0iPsKnZBPFCef3AZqR3VufM/YfzU+az95jsHKM8zLxnj3RmPx7ekmhhS60DaFBAB4iQr0vtaq73WuroAv+yUnK9nn3lwviqt0DU2IPF4QEaDfPZsL6LzryayqbAITyusyOaGDbM779xZXrTn40eG1W1a8BjVCeus04CqmVfvLVvNxKnzWPLxBgDq52W4PaVxLSLdgA+AKJAgtY+pWfCEpBbSQFaq8szOyPC6RXzz6ZKP10v3KybL5UOm8uXaQrw62ZiITxjapLoMLRl1cti0fhtnXDhBNhUW0f+clruPa5QzT4QWu3ZVfFyvXk6rep1GTqvcVzqcWNQ4UNr3xKuTzZdrC7l88FS6X/EASz7ewC/a1mP5kz24tFsTRIScTJOXwhZS41SitqUFWQ9AjFHTAeTTEafMmthTTjq+jpA835ERA8+Rwo8eFNk1S9zGGSLb/yCb358srY6vL4AbcmFbkXW3FomIBqYrpUQpNp7QuF4rgHDd00Ol9JWy7Sum2OEDuzultABy0nG5MnNMJ7FvXyryxdUyuFdzC0hOpulZC+NREzC7513TXtaMkooPbpQpI7tJgzoxASQ/N1MmjLpYytY9JVs+eCANXq7tfbKTNaPcFy9cXtmuRYPOSoFS6jmSB16bGuRFmwNq4EWdltUvyBEgbJAflSk3niYVCy8R+XCABAsvkfCDATKk9z9PgP0Lr2sXLr9JZMVwka9Gyc63bpDRV58uUT/ptdbNj5GWKfDXXXCyyF9ukfDTEWHjelkCfPfU4J/5APnZkZcA8YzaZLT+BpCIwY0e0EZ2vtJX5KMrRN7uJ+HCSyRYeInIh1f8XQJHfYBqjEacEBZX0bAgxn23nsmXs6/imj4ns+HbnXxTWMR1fU7m2fHnYqsDMYJ+8razqo87JvvmwU+tDAH2L+n33v9c0sqFVlpY504ceN4Jsvq5nuq+4e1omB8lLI4jTjA/QuMe7oKiph3Sk1LgGY0EFhcPad2sgOfv6cmg3m34est+Bv/qFEhYsA4iHl1/2ji4oVfLz7RWAihKE5kP3dZBt2ySnWjbLNfrekZjTcJii+NorZLKtJbJD6brkQmkm4WALq0ITH6dmkQUxihcdQDVAV3OOI4uv2gGZfHkWY6nlQ2t7fabV3PWfLvvTVl4/umq11vxIC7ixYSh/VtrQtGuNA5wWI9bJ2jhACn1A5Xnh6aQR/K4MRdwx53bssS6pA46pLFOLmi2IkFYXJWmh1KKRFVC7SuJAxSMmLopmvx9MqRBcQJbGRxoX9OcCNYK0Swf7SmKKxKOpB48rLN/iEApoEW4Cpjz25vf6GByo6HO8AmtQ2rF1WiFZw7tSimF5ymAMCNiBMBI0gO+Ud/bO4hAaAUdMZi8CCvX75W+I5fYucu2e0qhyyrtvsMBrc3KknTSx8Blzsk4pVT/yTNX9d+9vyoYfW1HaX5iPUVFgjARJiXvD3ggDYq/s19OH7F4EY2XGeHbb4vlvlnr3XNvbTHWiSG5Px+fwqRSGA/Y4SKQ0i+8DJwuIsMSgd329Btf++2vnKXGPbzM7i2txsvPRCkOnKAdjdnan60kC0N+lL1lCcY9scq2H/yOmrHwW2OdbI94+magQwrLIfcCRyKQNkNSFz0OtPd9dU9JRaJ0wrOfmXZXzuLxF1fauBNMXix54Ov+TrmoYaETRMDkRog7eHzOOttu8DtM+OPXpqQyKPV9dQ/QLhG6qRw8hTisHYlAejp5wJ4gkDHRKO09zzy7bXe5HTZ5qek48CX38sKvnYp6mOwI7jCJXtPioU0Kt2wfFTW8/PZfXcffvOOGTf3CbCuqtJ6nno1GaR8EMgbYw8GrpdrBOyoCcFDgKcCLx9kchvZ6z/M6e556c/WmPfqy2xfps4e8apd9stXp7Cgm5qfP9w90UlIZaMBk1s8MdXaEZSt3urNvWWIvG79Cr95Soj2t3vQ8OoehXB+Ps7kG8CMKt3/Eat9a9tJarUgNIv17tAxXz77Sydr/lXD5TbZp/WwBts29q18EYO7YTncM6nVCGpRorVZEjOlVo/9/2fXq0RBJD6Sjvr5GKTYCEvWNjBxwWvjX1waGzRrlCrAlLxptBjxEcj6LUmyI+vqamn38p4DXtprJlR3x9GhgFyB1cqKSmeGFQAmwnyTwPanf5PxAH/81q/HfCRr7vnqY5EVc+gI7AB4EGh+uzf8XO+RyPBKhDfCH1Kvmrb3hv/DvlB9jhxCpYf8W4P8HVeOD3yvYu0MAAAAASUVORK5CYII="""

def carpeta_programa():
    if getattr(sys, "frozen", False):
        return sys._MEIPASS
    return os.path.dirname(os.path.abspath(__file__))


@app.get("/logo.ico")
def logo():
    carpeta = os.path.join(carpeta_programa(), "icono")
    return send_from_directory(carpeta, "logo.ico")


@app.get("/favicon.png")
@app.get("/favicon.ico")
def favicon():
    from flask import Response
    respuesta = Response(base64.b64decode(FAVICON_ICO), mimetype="image/x-icon")
    respuesta.headers["Cache-Control"] = "no-cache"
    return respuesta


@app.route("/")
def index():
    editar = request.args.get("editar", type=int)
    if editar is not None and (editar < 0 or editar >= len(data[0])):
        editar = None
    return render_template_string(
        PAGE,
        topicos=topicos,
        filas=filas_vista(),
        dia=dia,
        fecha_pro=fecha_pro,
        editar=editar,
        meses=meses,
    )


@app.post("/agregar")
def agregar():
    boleta, error = leer_boleta()
    if error or boleta[0] not in topicos:
        flash(error or "Debe completar todos los campos.")
        return redirect(url_for("index"))

    for columna, valor in zip(data, boleta):
        columna.append(valor)
    return redirect(url_for("index"))


@app.post("/modificar/<int:indice>")
def modificar(indice):
    if indice < 0 or indice >= len(data[0]):
        flash("Esa boleta ya no está en la tabla.")
        return redirect(url_for("index"))

    boleta, error = leer_boleta()
    if error:
        flash(error)
        return redirect(url_for("index"))

    for columna, valor in zip(data, boleta):
        columna[indice] = valor
    return redirect(url_for("index"))


@app.post("/eliminar/<int:indice>")
def eliminar(indice):
    if indice < 0 or indice >= len(data[0]):
        flash("Esa boleta ya no está en la tabla.")
        return redirect(url_for("index"))
    for columna in data:
        columna.pop(indice)
    return redirect(url_for("index"))


@app.get("/guardar")
def guardar():
    nombre = f"Datos_guardados_no_procesados_{dia}.xlsx"
    return excel_descarga(workbook_sin_procesar(), nombre)


@app.post("/recuperar")
def recuperar():
    archivo = request.files.get("archivo")
    if not archivo or not archivo.filename:
        flash("Seleccione un archivo de recuperación.")
        return redirect(url_for("index"))

    try:
        excel = openpyxl.load_workbook(archivo)
        hoja2 = excel.active
        for i in range(1, hoja2.max_row + 1):
            topi = hoja2.cell(row=i, column=1).value
            prove = hoja2.cell(row=i, column=2).value
            nbole = hoja2.cell(row=i, column=3).value
            febole = hoja2.cell(row=i, column=4).value
            mon = hoja2.cell(row=i, column=5).value
            if not topi or not prove or not nbole or not febole or not mon:
                break
            data[0].append(topi)
            data[1].append(prove)
            data[2].append(nbole)
            iso = fecha_iso(febole)
            data[3].append(datetime.date.fromisoformat(iso) if iso else febole)
            data[4].append(int(mon))
        excel.close()
        flash("Datos recuperados con éxito.")
    except Exception as error:
        flash(f"No se pudo recuperar el archivo: {error}")
    return redirect(url_for("index"))


@app.post("/procesar")
def procesar():
    mes = request.form.get("mes", "").strip()
    if mes not in meses:
        flash("Seleccione el mes para procesar los datos.")
        return redirect(url_for("index"))
    nombre = f"DETALLE GASTOS BERFRE {mes} {dia}.xlsx"
    return excel_descarga(workbook_procesado(mes), nombre)


@app.post("/apagar")
def apagar():
    def cerrar():
        os._exit(0)

    threading.Timer(0.4, cerrar).start()
    return """
    <!doctype html>
    <html lang="es">
    <head><meta charset="utf-8"><title>Programa apagado</title></head>
    <body style="font-family: Segoe UI, sans-serif; margin: 24px; color: #1a1a1a;">
      <h1>Programa apagado</h1>
      <p>Ya puede cerrar esta pestaña.</p>
    </body>
    </html>
    """


if __name__ == "__main__":
    if os.environ.get("WERKZEUG_RUN_MAIN") == "true" or os.environ.get("FLASK_DEBUG") != "1":
        threading.Timer(1, lambda: webbrowser.open("http://127.0.0.1:5000/")).start()
    app.run(debug=os.environ.get("FLASK_DEBUG") == "1", port=5000)
