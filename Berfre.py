import os
import os.path as path
from pathlib import Path
import sys
import io
from matplotlib.pylab import rint
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill, NamedStyle
import tkinter as tk
from tkinter import filedialog
import datetime
from sympy import true
#from diccionariotemp import ubicaciones, precios, material
import private.informacion_delicada.diccionario

# inicializa variable de tiempo
fecha = datetime.date.today()
dia = str(fecha.year) + str(fecha.month) + str(fecha.day)
print("Codigo de dia: ", dia)

#
#PRUEBA DE FLASK
#

def maquillaje():
    bordes = Border(
            bottom=Side(border_style="medium", color="000000"),
            left=Side(border_style="thin"),
            right=Side(border_style="thin"),
            top=Side(border_style="thin")
            )
    color = PatternFill(
            start_color='a9e5e5', end_color='5CB800', fill_type='solid')
    vacio = PatternFill(
            start_color='f9e37c', end_color='5CB800', fill_type='solid')
    return bordes, color, vacio

#Funcion para realizar filtros y mostrar solo el stock de M501
def filtro1(ruta):
    # inicializa manejo de archivos
    #test = input("Ingrese el nombre del archivo de recuperacion con su extencion ==> ")
    excel = openpyxl.load_workbook(ruta)
    mango = openpyxl.Workbook()
    hoja2 = excel.active
    hoja = mango.active

    #Variable que maneja la fila en la que se esta escribiendo
    x = 4

    #Diseño de las celdas
    bordes, color, vacio = maquillaje()

    hoja.column_dimensions['A'].width = 20
    hoja.column_dimensions['B'].width = 52
    hoja.column_dimensions['C'].width = 20
    hoja['A3'].border = bordes
    hoja['B3'].border = bordes
    hoja['C3'].border = bordes
    hoja['A1'].border = bordes
    hoja['B1'].border = bordes
    hoja['A3'].font = Font(bold=True)
    hoja['B3'].font = Font(bold=True)
    hoja['C3'].font = Font(bold=True)
    hoja['A1'].font = Font(bold=True, size=12)
    hoja['B1'].font = Font(bold=True, size=12)
    hoja['A3'].fill = color
    hoja['B3'].fill = color
    hoja['C3'].fill = color
    hoja['A1'].fill = color
    hoja['B1'].fill = color

    #Titulo del documento
    hoja['A1'] = "Almacen"
    hoja['B1'] = "M501"

    #Titlos de las columnas
    hoja['A3'] = "Etiqueta de fila"
    hoja['B3'] = "Descripcion del producto"
    hoja['C3'] = "Libre utilización"

    #Filtros para el archivo
    hoja.auto_filter.ref = "A3:C3"
    hoja.auto_filter.add_sort_condition("C4:C" + str(hoja.max_row))

    #Manejo de archivos
    for i in range(1, hoja2.max_row + 1):
        bodega = hoja2.cell(row = i, column = 3).value
        if bodega == "M501":
            if hoja2.cell(row = i, column = 2).value != "NULO":
                id = hoja2.cell(row = i, column = 1).value
                hoja.cell(row = x, column = 1, value = id).border = bordes
                descripcion = hoja2.cell(row = i, column = 2).value
                hoja.cell(row = x, column = 2, value = descripcion).border = bordes
                utilizacion = hoja2.cell(row = i, column = 4).value
                hoja.cell(row = x, column = 3, value = utilizacion).border = bordes
                x += 1

    #hoja['A1'] = hoja2.cell(row = 1, column = 1).value
    #hoja['B3'] = 'EMPRESA: PRETORIANOS SEGURIDAD'
    print(f"i = {i} y x = {x}")
    nombre_archivo = "Prueba_de_planilla_M501.xlsx"
    mango.save(nombre_archivo)
    ruta_creacion = path.abspath(nombre_archivo)
    return nombre_archivo, ruta_creacion

# Versión preview de filtro1: retorna columnas y filas como listas (sin guardar archivo)
def filtro1_preview(ruta):
    excel = openpyxl.load_workbook(ruta)
    hoja2 = excel.active
    columnas = ["Etiqueta de fila", "Descripcion del producto", "Libre utilización"]
    filas = []
    for i in range(1, hoja2.max_row + 1):
        bodega = hoja2.cell(row=i, column=3).value
        if bodega == "M501":
            if hoja2.cell(row=i, column=2).value != "NULO":
                filas.append([
                    hoja2.cell(row=i, column=1).value,
                    hoja2.cell(row=i, column=2).value,
                    hoja2.cell(row=i, column=4).value,
                ])
    return columnas, filas

#Planilla todas la bodegas
def total(ruta):
    # inicializa manejo de archivos
    excel = openpyxl.load_workbook(ruta)
    mango = openpyxl.Workbook()
    hoja2 = excel.active
    hoja = mango.active

    #Variable que maneja la fila en la que se esta escribiendo
    x = 4

    #Diseño de las celdas
    bordes, color, vacio = maquillaje()

    hoja.column_dimensions['A'].width = 20
    hoja.column_dimensions['B'].width = 52
    hoja.column_dimensions['C'].width = 15
    hoja.column_dimensions['D'].width = 15
    hoja.column_dimensions['E'].width = 15
    hoja.column_dimensions['F'].width = 15
    hoja.column_dimensions['G'].width = 15

    #Titulo del documento
    hoja['A1'] = "Stock por bodega"
    
    #Titlos de las columnas
    hoja['A3'] = "Material"
    hoja['B3'] = "Descripcion del producto"
    hoja['C3'] = "M501"
    hoja['D3'] = "M502"
    hoja['E3'] = "M503"
    hoja['F3'] = "M504"
    hoja['G3'] = "M505"
    hoja['H3'] = "Total"

    #Diseño
    hoja['A3'].font = Font(bold=True)
    hoja['B3'].font = Font(bold=True)
    hoja['C3'].font = Font(bold=True)
    hoja['D3'].font = Font(bold=True)
    hoja['E3'].font = Font(bold=True)
    hoja['F3'].font = Font(bold=True)
    hoja['G3'].font = Font(bold=True)
    hoja['H3'].font = Font(bold=True)
    hoja['A1'].font = Font(bold=True)
    hoja['A3'].fill = color
    hoja['B3'].fill = color
    hoja['C3'].fill = color
    hoja['D3'].fill = color
    hoja['E3'].fill = color
    hoja['F3'].fill = color
    hoja['G3'].fill = color
    hoja['H3'].fill = color
    hoja['A1'].fill = color
    hoja['A3'].border = bordes
    hoja['B3'].border = bordes
    hoja['C3'].border = bordes
    hoja['D3'].border = bordes
    hoja['E3'].border = bordes
    hoja['F3'].border = bordes
    hoja['G3'].border = bordes
    hoja['H3'].border = bordes
    hoja['A1'].border = bordes

    #Filtros para el archivo
    hoja.auto_filter.ref = "A3:H3"

    for i in range(2, hoja2.max_row + 1):
        if hoja2.cell(row = i, column = 2).value != "NULO":
            hoja.cell(row = x, column = 3).border = bordes
            hoja.cell(row = x, column = 4).border = bordes
            hoja.cell(row = x, column = 5).border = bordes
            hoja.cell(row = x, column = 6).border = bordes
            hoja.cell(row = x, column = 7).border = bordes
            bodega = hoja2.cell(row = i, column = 1).value

            if bodega != hoja2.cell(row = i - 1, column = 1).value:
                id = hoja2.cell(row = i, column = 1).value
                hoja.cell(row = x, column = 1, value = id).border = bordes
                descripcion = hoja2.cell(row = i, column = 2).value
                hoja.cell(row = x, column = 2, value = descripcion).border = bordes
                M501 = hoja2.cell(row = i, column = 4).value
                hoja.cell(row = x, column = 3, value = M501)

                if hoja2.cell(row = i + 1, column = 3).value == "M502":
                    M502 = hoja2.cell(row = i + 1, column = 4).value
                    hoja.cell(row = x, column = 4, value = M502)

                if hoja2.cell(row = i + 2, column = 3).value == "M503":
                    M503 = hoja2.cell(row = i + 2, column = 4).value
                    hoja.cell(row = x, column = 5, value = M503)

                if hoja2.cell(row = i + 3, column = 3).value == "M504":
                    M504 = hoja2.cell(row = i + 3, column = 4).value
                    hoja.cell(row = x, column = 6, value = M504)

                if hoja2.cell(row = i + 4, column = 3).value == "M505":
                    M505 = hoja2.cell(row = i + 4, column = 4).value
                    hoja.cell(row = x, column = 7, value = M505)

                hoja.cell(row = x, column = 8, value = M501 + M502 + M503 + M504 + M505).font = Font(bold=True)
                hoja.cell(row = x, column = 8).border = bordes
                i = i + 4
                # Después de escribir los valores de las bodegas en la fila x:
                for col_letra in ['C', 'D', 'E', 'F', 'G']:
                    val = hoja[f'{col_letra}{x}'].value
                    if val is None or val == "":
                        hoja[f'{col_letra}{x}'].fill = vacio
                x += 1

    print(f"i = {i} y x = {x}")
    nombre_archivo = "Prueba_de_planilla_stock_region.xlsx"
    mango.save(nombre_archivo)
    ruta_creacion = path.abspath(nombre_archivo)
    return nombre_archivo, ruta_creacion

# Versión preview de total: retorna columnas y filas como listas (sin guardar archivo)
def total_preview(ruta):
    excel = openpyxl.load_workbook(ruta)
    hoja2 = excel.active
    columnas = ["Material", "Descripcion del producto", "M501", "M502", "M503", "M504", "M505", "Total"]
    filas = []
    for i in range(2, hoja2.max_row + 1):
        if hoja2.cell(row=i, column=2).value != "NULO":
            bodega = hoja2.cell(row=i, column=1).value
            if bodega != hoja2.cell(row=i - 1, column=1).value:
                M501 = hoja2.cell(row=i, column=4).value or 0
                M502 = hoja2.cell(row=i+1, column=4).value if hoja2.cell(row=i+1, column=3).value == "M502" else 0
                M503 = hoja2.cell(row=i+2, column=4).value if hoja2.cell(row=i+2, column=3).value == "M503" else 0
                M504 = hoja2.cell(row=i+3, column=4).value if hoja2.cell(row=i+3, column=3).value == "M504" else 0
                M505 = hoja2.cell(row=i+4, column=4).value if hoja2.cell(row=i+4, column=3).value == "M505" else 0
                filas.append([
                    hoja2.cell(row=i, column=1).value,
                    hoja2.cell(row=i, column=2).value,
                    M501, M502, M503, M504, M505,
                    M501 + M502 + M503 + M504 + M505,
                ])
    return columnas, filas

#Inventario de la bodega M501 con ubicacion
def inventario(ruta, dicub):
    # inicializa manejo de archivos
    excel = openpyxl.load_workbook(ruta)
    mango = openpyxl.Workbook()
    hoja2 = excel.active
    hoja = mango.active

    #Variable que maneja la fila en la que se esta escribiendo
    x = 4

    #Diseño de las celdas
    bordes, color, vacio = maquillaje()

    hoja.column_dimensions['A'].width = 20
    hoja.column_dimensions['B'].width = 52
    hoja.column_dimensions['C'].width = 20
    hoja.column_dimensions['D'].width = 15
    hoja.column_dimensions['E'].width = 15
    hoja['A3'].border = bordes
    hoja['B3'].border = bordes
    hoja['C3'].border = bordes
    hoja['D3'].border = bordes
    hoja['E3'].border = bordes
    hoja['A1'].border = bordes
    hoja['B1'].border = bordes
    hoja['A3'].font = Font(bold=True)
    hoja['B3'].font = Font(bold=True)
    hoja['C3'].font = Font(bold=True)
    hoja['D3'].font = Font(bold=True)
    hoja['E3'].font = Font(bold=True)
    hoja['A1'].font = Font(bold=True, size=12)
    hoja['B1'].font = Font(bold=True, size=12)
    hoja['A3'].fill = color
    hoja['B3'].fill = color
    hoja['C3'].fill = color
    hoja['D3'].fill = color
    hoja['A1'].fill = color
    hoja['B1'].fill = color

    #Titulo del documento
    hoja['A1'] = "Almacen"
    hoja['B1'] = "M501"

    #Titlos de las columnas
    hoja['A3'] = "Etiqueta de fila"
    hoja['B3'] = "Descripcion del producto"
    hoja['C3'] = "Ubicacion"
    hoja['D3'] = "Libre utilización"
    hoja['E3'] = "Existencia"

    #Ordena la planilla por ubicacion
    hoja.auto_filter.ref = "A3:E3"
    hoja.auto_filter.add_sort_condition("C4:C" + str(hoja.max_row))

    #Manejo de archivos
    encabezados = ["Etiqueta de fila", "Descripcion del producto", "Ubicacion", "Libre utilización", "Existencia"]
    for col_idx, texto in enumerate(encabezados, start=1):
        cell = hoja.cell(row=3, column=col_idx, value=texto)
        cell.font = Font(bold=True)
        cell.border = bordes

    # 3. Lectura y Filtrado en memoria
    filas_filtradas = []

    # iter_rows es infinitamente más rápido que recorrer celda por celda
    for row in hoja2.iter_rows(values_only=True):
        if not row or len(row) < 4:
            continue
            
        id_prod = row[0]          # Columna 1
        descripcion = row[1]      # Columna 2
        bodega = row[2]           # Columna 3
        utilizacion = row[3]      # Columna 4

        if bodega == "M501" and descripcion != "NULO":
            # Buscar ubicación en el diccionario
            id_str = str(id_prod)
            ubicacion = ""
            if id_str in dicub and dicub[id_str]:
                ubicacion = dicub[id_str]

            filas_filtradas.append({
                'id': id_prod,
                'descripcion': descripcion,
                'ubicacion': ubicacion,
                'utilizacion': utilizacion,
                'existencia': None
            })

    # 4. ORDENAMIENTO EN MEMORIA (Clave de la velocidad)
    # Criterio: Extrae los 2 primeros caracteres si existen números, de lo contrario da prioridad 0 o valor por defecto
    # 4. ORDENAMIENTO EN MEMORIA
    def obtener_clave_ordenamiento(item):
        ub = str(item['ubicacion']).strip() if item['ubicacion'] else ""
        
        # Si NO tiene ubicación (vacía o None)
        if not ub:
            # Devuelve (1, 0): El '1' la manda al final
            return (1, 0)
        
        # Si SÍ tiene ubicación y empieza con al menos 2 dígitos
        if len(ub) >= 2 and ub[:2].isdigit():
            # Devuelve (0, número): El '0' la pone arriba, ordenada por su valor numérico
            return (0, int(ub[:2]))
        
        # Si tiene texto pero no empieza con números (caso borde)
        return (0, 999)

    # Ordenamos la lista en Python usando la tupla como prioridad
    filas_filtradas.sort(key=obtener_clave_ordenamiento)

    # 5. Escritura rápida en la hoja de destino
    x = 4
    for item in filas_filtradas:
        hoja.cell(row=x, column=1, value=item['id']).border = bordes
        hoja.cell(row=x, column=2, value=item['descripcion']).border = bordes
        hoja.cell(row=x, column=3, value=item['ubicacion']).border = bordes
        if hoja.cell(row=x, column= 3).value == "":
            hoja[f'C{x}'].fill = vacio
        hoja.cell(row=x, column=4, value=item['utilizacion']).border = bordes
        hoja.cell(row=x, column=5, value=item['existencia']).border = bordes
        x += 1

    # Agregar Autofiltro visual
    hoja.auto_filter.ref = f"A3:E{x-1}"

    # Guardar archivo
    nombre_archivo = "Prueba_de_planilla_invetario.xlsx"
    mango.save(nombre_archivo)
    ruta_creacion = path.abspath(nombre_archivo)
    print("¡Proceso completado con éxito!")
    return nombre_archivo, ruta_creacion

# Versión preview de inventario: retorna columnas y filas como listas (sin guardar archivo)
def inventario_preview(ruta, dicub):
    excel = openpyxl.load_workbook(ruta)
    hoja2 = excel.active
    columnas = ["Etiqueta de fila", "Descripcion del producto", "Ubicacion", "Libre utilización", "Existencia"]
    filas_filtradas = []

    for row in hoja2.iter_rows(values_only=True):
        if not row or len(row) < 4:
            continue
        id_prod, descripcion, bodega, utilizacion = row[0], row[1], row[2], row[3]
        if bodega == "M501" and descripcion != "NULO":
            id_str = str(id_prod)
            ubicacion = dicub.get(id_str, "") if id_str in dicub else ""
            filas_filtradas.append({
                'id': id_prod,
                'descripcion': descripcion,
                'ubicacion': ubicacion,
                'utilizacion': utilizacion,
                'existencia': None,
            })

    def obtener_clave_ordenamiento(item):
        ub = str(item['ubicacion']).strip() if item['ubicacion'] else ""
        if not ub:
            return (1, 0)
        if len(ub) >= 2 and ub[:2].isdigit():
            return (0, int(ub[:2]))
        return (0, 999)

    filas_filtradas.sort(key=obtener_clave_ordenamiento)
    filas = [[f['id'], f['descripcion'], f['ubicacion'], f['utilizacion'], f['existencia']]
             for f in filas_filtradas]
    return columnas, filas

#Funcion para editar diccionarios, añadiendo, eliminando y editando datos
def modificar_diccionarios(dic, name):
    imp = int(input("Elije el numero de la opcion que quieres usar\n1.- Agregar un valor al diccionario\n2.- Modificar un valor del diccionario\n3.- Eliminar un valor del diccionario\n4.- Importar diccionario\n5.- Exportar diccionario\n>> "))
    if imp == 1:
        clave = input("Ingrese la clave que quiere agregar ==> ")
        valor = input("Ingrese el valor que quiere agregar ==> ")
        dic[clave] = valor
    elif imp == 2:
        clave = input("Ingrese la clave que quiere modificar ==> ")
        if clave in dic:
            valor = input("Ingrese el nuevo valor ==> ")
            dic[clave] = valor
        else:
            print("La clave no existe en el diccionario.")
    elif imp == 3:
        clave = input("Ingrese la clave que quiere eliminar ==> ")
        if clave in dic:
            del dic[clave]
        else:
            print("La clave no existe en el diccionario.")
    elif imp == 4:
        archivo = pedir_archivo_visual()
        try:
            excel = openpyxl.load_workbook(archivo)
            hoja = excel.active
            for i in range(2, hoja.max_row + 1):
                clave = str(hoja.cell(row=i, column=1).value)
                valor = str(hoja.cell(row=i, column=2).value)
                if clave is not None and valor is not None:
                    dic[clave] = valor
        except Exception as e:
            print(f"Error al importar el diccionario: {e}")
    elif imp == 5:
        j = 2

        excel = openpyxl.Workbook()
        hoja = excel.active
        hoja['A1'] = "Codigo"
        hoja['B1'] = name
        for i in dic.keys():
            clave = i
            valor = dic[i]
            
            hoja.cell(row=j, column=1).value = clave
            hoja.cell(row=j, column=2).value = valor
            j += 1
        excel.save(f"diccionario_{name}_exportado.xlsx")
    return

#Añadir ventas al excel de ventas
def añadir_venta(dic_mat, dic_ub, dic_pre, dic_sto, ruta):
    export = openpyxl.load_workbook(ruta)
    temp = export.active
    #diccionarios para el stock
    M501 = {}
    for p in  range(2, temp.max_row + 1):
        if temp.cell(row = p, column = 3).value == "M501":
            M501[str(temp.cell(row = p, column = 1).value)] = temp.cell(row = p, column = 4).value
    M502 = {}
    for p in  range(2, temp.max_row + 1):
        if temp.cell(row = p, column = 3).value == "M502":
            M502[str(temp.cell(row = p, column = 1).value)] = temp.cell(row = p, column = 4).value
    M503 = {}
    for p in  range(2, temp.max_row + 1):
        if temp.cell(row = p, column = 3).value == "M503":
            M503[str(temp.cell(row = p, column = 1).value)] = temp.cell(row = p, column = 4).value
    M504 = {}
    for p in  range(2, temp.max_row + 1):
        if temp.cell(row = p, column = 3).value == "M504":
            M504[str(temp.cell(row = p, column = 1).value)] = temp.cell(row = p, column = 4).value
    M505 = {}
    for p in  range(2, temp.max_row + 1):
        if temp.cell(row = p, column = 3).value == "M505":
            M505[str(temp.cell(row = p, column = 1).value)] = temp.cell(row = p, column = 4).value
    cont = 10
    print("Elija la orden de venta que quiera cargar ")
    ruta_orden = pedir_archivo_visual()
    excel = openpyxl.load_workbook(ruta_orden)
    op = int(input("¿Quiere cargar una pagina de ventas o prefiere crear una nueva?\n1.- Cargar archivo xlsx existente\n2.- Crear nuevo archivo de ventas\n>> "))
    if op < 1 or op > 2:
        print("Una pega...")
        return
    if op == 1:
        ruta_arch = pedir_archivo_visual()
        mango = openpyxl.load_workbook(ruta_arch)
        hoja = mango.active
    else:
        mango = openpyxl.Workbook()
        hoja = mango.active
        last_pos = 1
        encabezado_ventas(hoja, last_pos)
    hoja2 = excel.active
    for i in range(1, hoja.max_row + 1):
        if hoja.cell(row = i, column = 1).value == "Pos":
            #print ("new")
            last_pos = i
    for j in range(2, hoja2.max_row + 1):
        #ID
        hoja.cell(row = last_pos + 1, column = 1, value = cont)
        #Codigo de material
        hoja.cell(row = last_pos + 1, column = 2, value = hoja2.cell(row = j, column = 1).value)
        #Descripcion del material
        if str(hoja2.cell(row = j, column = 1).value) in dic_mat:
            hoja.cell(row = last_pos + 1, column = 3, value = dic_mat[str(hoja2.cell(row = j, column = 1).value)])
        #Ubicacion del material
        if str(hoja2.cell(row = j, column = 1).value) in dic_ub:
            hoja.cell(row = last_pos + 1, column = 4, value = dic_ub[str(hoja2.cell(row = j, column = 1).value)])
        #stock M501
        if str(hoja2.cell(row = j, column = 1)) in dic_sto:
            hoja.cell(row = last_pos + 1, column = 5, value = (dic_sto[str(hoja2.cell(row = j, column = 1))]))
        else:
            hoja.cell(row = last_pos + 1, column = 5, value = 0)
        #Stock M502

        #Stock M503

        #Stock M504

        #Stock M505

        #Cantidad
        cant = hoja2.cell(row = j, column = 2).value
        hoja.cell(row = last_pos + 1, column = 7, value = cant)
        #Entregar
        if cant >= int(dic_sto[str(hoja2.cell(row = j, column = 1).value)]):
            hoja.cell(row = last_pos + 1, column = 8, value = (cant - int(dic_sto[str(hoja2.cell(row = j, column = 1).value)])))
            infra = 1
        else:
            hoja.cell(row = last_pos + 1, column = 8, value = 0)
            infra = 0
        #Peso y tiras
        #Columna 9 y 10 se añaden datos manualmente, se rellena con 0 mientras tanto
        hoja.cell(row = last_pos + 1, column = 9, value = 0)
        hoja.cell(row = last_pos + 1, column = 10, value = 0)
        #Diferencia
        hoja.cell(row = last_pos + 1, column = 11, value = (cant - int(hoja.cell(row = last_pos + 1, column = 8))))
        #Verdadero o falso
        if hoja.cell(row = last_pos + 1, column = 8) == hoja.cell(row = last_pos + 1, column = 7):
            hoja.cell(row = last_pos + 1, column = 12, value = "VERDADERO")
        else:
            hoja.cell(row = last_pos + 1, column = 12, value = "FALSO")
        #Kg por unidad y peso total
        #Columnas 13 y 14 se rellenan manualmente por ahora, se rellena con 0
        hoja.cell(row = last_pos + 1, column = 13, value = 0)
        hoja.cell(row = last_pos + 1, column = 14, value = 0)
        #Precio
        if str(hoja2.cell(row = j, column = 1).value) in dic_pre:
            precio = dic_pre[str(hoja2.cell(row = j, column = 1).value)]
            hoja.cell(row = last_pos + 1, column = 15, value = precio)
        else:
            precio = 0
        #Precio total
        hoja.cell(row = last_pos + 1, column = 16, value = (cant * precio))
        last_pos += 1
        cont += 10
    encabezado_ventas(hoja, last_pos + 1)
    mango.save(f"prueba_venta.xlsx")
    return

def encabezado_ventas(hoja, last_pos):
    amarillo = ["Pos", "Codigo", "Descripcion", "Ubicación", "M501", "M502", "M503", "M504", "M505", "Solicitado", "Entregar", "Med", "Tiras", "Dif", "Comp", "kg x U", "Kg Total GD", "$ x U", "$ Total GD", "Movimiento", "Almacen", "N°Venta", "Cant:GD", "GD Esval",	"Estado", "Fecha", "Dif.Pend",	"GD Esval",	"Fecha", "Observacion"]
    gris = ["Pos",	"Codigo",	"SOLICITUD",	"COMPARA CODIGO",	"COMPARA CANT"]

    for i in range(0, len(amarillo)):
        hoja.cell(row = last_pos, column = i + 1, value = amarillo[i]).fill = PatternFill(
                                                        start_color='ffff00', end_color='5CB800', fill_type='solid')
    i += 2
    for j in range(0, len(gris)):
        hoja.cell(row = last_pos, column = i, value = gris[j])
        i += 1
    return

# ─── HELPERS PARA FLASK (sin input()) ────────────────────────────────────────

def dic_agregar(dic: dict, clave: str, valor: str) -> dict:
    """Agrega o sobreescribe una clave en el diccionario."""
    dic[clave] = valor
    return dic

def dic_modificar(dic: dict, clave: str, valor: str):
    """Modifica el valor de una clave existente. Retorna el dict o None si la clave no existe."""
    if clave not in dic:
        return None
    dic[clave] = valor
    return dic

def dic_eliminar(dic: dict, clave: str):
    """Elimina una clave. Retorna el dict o None si la clave no existe."""
    if clave not in dic:
        return None
    del dic[clave]
    return dic

def dic_exportar_bytes(dic: dict, name: str) -> io.BytesIO:
    """Exporta el diccionario a un archivo Excel en memoria (BytesIO)."""
    excel = openpyxl.Workbook()
    hoja = excel.active
    hoja['A1'] = "Codigo"
    hoja['B1'] = name
    for j, (clave, valor) in enumerate(dic.items(), start=2):
        hoja.cell(row=j, column=1, value=clave)
        hoja.cell(row=j, column=2, value=valor)
    buffer = io.BytesIO()
    excel.save(buffer)
    buffer.seek(0)
    return buffer

def dic_importar_excel(file_bytes: bytes) -> dict:
    """Lee un xlsx (bytes) y retorna un dict {clave: valor} desde la fila 2 en adelante."""
    buffer = io.BytesIO(file_bytes)
    excel = openpyxl.load_workbook(buffer)
    hoja = excel.active
    resultado = {}
    for i in range(2, hoja.max_row + 1):
        clave = hoja.cell(row=i, column=1).value
        valor = hoja.cell(row=i, column=2).value
        if clave is not None:
            resultado[str(clave)] = str(valor) if valor is not None else ""
    return resultado

#Seleccion de archivo visualmente
def pedir_archivo_visual():
    # 1. Crear una ventana raíz oculta para que no aparezca una ventana vacía de fondo
    root = tk.Tk()
    root.withdraw()
    
    # 2. Forzar a que la ventana de selección aparezca al frente de todo
    root.attributes('-topmost', True)
    
    # 3. Abrir el cuadro de diálogo para seleccionar el archivo
    ruta_archivo = filedialog.askopenfilename(
        title="Selecciona un archivo",
        filetypes=[("Todos los archivos", "*.*"), ("Archivos de Texto", "*.txt"), ("Documentos PDF", "*.pdf")]
    )
    
    # 4. Destruir la ventana raíz al terminar
    root.destroy()
    
    return ruta_archivo

#Selecciona que filtro se quiere usar y ejecuta la funcion correspondiente
def main():
    # Ejecución del ejemplo
    ruta_seleccionada = pedir_archivo_visual()

    if ruta_seleccionada:
        print(f"\nRuta seleccionada visualmente: {ruta_seleccionada}")
    else:
        print("\nEl usuario canceló la selección.")
        return

    #Se inicializan los diccionarios
    dicub = private.informacion_delicada.diccionario.ubicaciones
    dicpre = private.informacion_delicada.diccionario.precios
    dicmat = private.informacion_delicada.diccionario.creardicmar(ruta_seleccionada)
    dicsto = private.informacion_delicada.diccionario.creardictotal(ruta_seleccionada)
    dicalm = private.informacion_delicada.diccionario.almacenes
    diccrit = private.informacion_delicada.diccionario.critico
    diccom = private.informacion_delicada.diccionario.comprador

    print("RECUERDA QUE ESTAS TRABAJANDO EN LA LINEA 531")

    z = int(input("Elije el numero de la opcion que quieres usar\n1.- Filtro para solo M501\n2.- Filtro stock total\n3.- Planilla de inventario\n4.- Modificar diccionarios\n5.- Añadir una venta\n>> "))

    # Confirmación antes de generar el archivo
    confirmar = input(f"¿Seguro que quieres generar el archivo para la opción {z}? (s/n): ").strip().lower()
    if confirmar != 's':
        print("Operación cancelada.")
        return

    if z == 1:
        nombre, ruta_out = filtro1(ruta_seleccionada)
        print(f"Archivo generado: {nombre}\nUbicación: {ruta_out}")
    if z == 2:
        nombre, ruta_out = total(ruta_seleccionada)
        print(f"Archivo generado: {nombre}\nUbicación: {ruta_out}")
    if z == 3:
        nombre, ruta_out = inventario(ruta_seleccionada, dicub)
        print(f"Archivo generado: {nombre}\nUbicación: {ruta_out}")
    if z == 4:
        zan = int(input("Elije el numero de la opcion que quieres usar\n1.- Modificar diccionario de ubicaciones\n2.- Modificar diccionario de precios\n>> "))
        if zan == 1:
            modificar_diccionarios(dicub, "ubicaciones")
        if zan == 2:
            modificar_diccionarios(dicpre, "precios")
    if z == 5:
        añadir_venta(dicmat, dicub, dicpre, dicsto, ruta_seleccionada)

if __name__ == '__main__':
    print("Esto solo se ejecutará si corres Berfre.py directamente")
    main()