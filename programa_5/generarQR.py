import qrcode
import os
import openpyxl
import tkinter as tk
from tkinter import filedialog

carpeta = r".\programa_5\QRmateriales"   # carpeta destino
os.makedirs(carpeta, exist_ok=True)


# 1. Crear una ventana raíz oculta para que no aparezca una ventana vacía de fondo
root = tk.Tk()
root.withdraw()
# 2. Forzar a que la ventana de selección aparezca al frente de todo
root.attributes('-topmost', True)
# 3. Abrir el cuadro de diálogo para seleccionar el archivo
ruta_archivo = filedialog.askopenfilename(
    title="Selecciona un archivo",
    filetypes=[("Todos los archivos", "*.*"), ("Archivos de Texto", "*.txt"), ("Documentos PDF", "*.pdf")]
)
# 4. Destruir la ventana raíz al terminar
root.destroy()
#return ruta_archivo

excel = openpyxl.load_workbook(ruta_archivo)
hoja2 = excel.active

materiales = []

for i in range(1, hoja2.max_row + 1):
    bodega = hoja2.cell(row = i, column = 3).value
    if bodega == "M501":
        if hoja2.cell(row = i, column = 2).value != "NULO":
            id = hoja2.cell(row = i, column = 1).value
            materiales.append(id)
print(len(materiales))
excel.close()

qr = qrcode.QRCode(
    version=2,
    box_size=10,
    border=5,
    error_correction=qrcode.constants.ERROR_CORRECT_M
)

for material in materiales:
    qr.clear()  # Limpiar el contenido del QR antes de agregar nuevos datos
    url = f"https://practica-2.anibalmunoz74.workers.dev/{material}"

    qr.add_data(url)
    qr.make(fit=True) 

    img = qr.make_image(fill_color='black', back_color='white')
    img.save(os.path.join(carpeta, f"qr_{material}.png"))