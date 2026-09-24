precios = {

}

ubicaciones = {
    20000000: "Ubicación 1",
    20000001: "Ubicación 2",
    20000002: "Ubicación 3",
    20000003: "Ubicación 4",
    20000004: "Ubicación 5",
    20000005: "Ubicación 6",
    20000006: "Ubicación 7",
    20000007: "Ubicación 8",
    20000008: "Ubicación 9",
    20000009: "Ubicación 10",
    20000010: "Ubicación 11",
    20000011: "Ubicación 12",
    20000012: "Ubicación 13",
    20000013: "Ubicación 14",
    20000014: "Ubicación 15",
    20000015: "Ubicación 16",
    20000040: "Ubicación 17",
    20000041: "Ubicación 1",
}

material = {
    
}

almacenes = {

}

comprador = {

}

critico = {

}

total = {

}

def creardictotal(test):
    import openpyxl
    excel = openpyxl.load_workbook(test)
    hoja = excel.active
    for i in range(2, hoja.max_row + 1):
        if str(hoja.cell(row=i, column=3).value).strip() == "M501" and hoja.cell(row=i, column=2).value != "NULO":
            codigo = hoja.cell(row=i, column=1).value
            dato = hoja.cell(row=i, column=4).value
            if codigo is not None:
                total[str(codigo)] = str(dato) if dato is not None else ""
    return total