from Berfre import (dia, filtro1, filtro1_preview,
                    inventario, inventario_preview,
                    total as berfre_total, total_preview,
                    pedir_archivo_visual,
                    modificar_diccionarios,
                    dic_agregar, dic_modificar, dic_eliminar,
                    dic_exportar_bytes, dic_importar_excel)
import private.informacion_delicada.diccionario
import io
import json
import os
import openpyxl
from flask import (Flask, render_template, request, redirect,
                   url_for, session, send_file)
import sys

# Detectar si se está ejecutando como un .exe compilado
if getattr(sys, 'frozen', False):
    template_folder = os.path.join(sys._MEIPASS, 'templates')
    static_folder = os.path.join(sys._MEIPASS, 'static')
    front_flask = Flask(__name__, template_folder=template_folder, static_folder=static_folder)
else:
    front_flask = Flask(__name__)

# Clave secreta requerida para que Flask pueda usar 'session'
front_flask.secret_key = 'practica2_clave_secreta_desarrollo'

# ─── PERSISTENCIA Y SINCRONIZACIÓN DE DICCIONARIOS ───────────────────────────

_DIR_INFO = os.path.join(os.path.dirname(__file__), 'private', 'informacion_delicada')
_RUTA_DICCIONARIO_PY = os.path.join(_DIR_INFO, 'diccionario.py')
_JSON_FILES = {
    'ubicaciones': os.path.join(_DIR_INFO, 'ubicaciones_override.json'),
    'precios':     os.path.join(_DIR_INFO, 'precios_override.json'),
    'material':    os.path.join(_DIR_INFO, 'material_override.json'),
    'almacenes':   os.path.join(_DIR_INFO, 'almacenes_override.json'),
    'comprador':   os.path.join(_DIR_INFO, 'comprador_override.json'),
    'critico':     os.path.join(_DIR_INFO, 'critico_override.json'),
    'total':       os.path.join(_DIR_INFO, 'total_override.json'),
}

_DIC_LABELS = {
    'ubicaciones': 'Ubicaciones',
    'precios':     'Precios',
    'material':    'Materiales',
    'almacenes':   'Almacenes',
    'comprador':   'Comprador',
    'critico':     'Crítico',
    'total':       'Total',
}
_DIC_VALIDOS = list(_DIC_LABELS.keys())


def _cargar_total_desde_excel(dic_dest: dict, ruta_archivo: str) -> dict:
    """
    Carga el diccionario total desde un Excel con la lógica:
    Clave = Columna 1, Dato = Columna 4, solo si Columna 3 == 'M501' y Columna 2 != 'NULO'.
    """
    excel = openpyxl.load_workbook(ruta_archivo)
    hoja = excel.active
    for i in range(2, hoja.max_row + 1):
        c3 = str(hoja.cell(row=i, column=3).value).strip() if hoja.cell(row=i, column=3).value is not None else ""
        c2 = hoja.cell(row=i, column=2).value
        if c3 == "M501" and c2 != "NULO":
            codigo = hoja.cell(row=i, column=1).value
            dato = hoja.cell(row=i, column=4).value
            if codigo is not None:
                dic_dest[str(codigo)] = str(dato) if dato is not None else ""
    return dic_dest


def _guardar_en_diccionario_py():
    """
    Reescribe private/informacion_delicada/diccionario.py con los diccionarios vivos
    para que tanto Flask como Berfre.py o cualquier otra función utilicen los
    datos modificados permanentemente.
    """
    try:
        mat = getattr(private.informacion_delicada.diccionario, 'material', {})
        pre = getattr(private.informacion_delicada.diccionario, 'precios', {})
        ubi = getattr(private.informacion_delicada.diccionario, 'ubicaciones', {})
        alm = getattr(private.informacion_delicada.diccionario, 'almacenes', {})
        com = getattr(private.informacion_delicada.diccionario, 'comprador', {})
        cri = getattr(private.informacion_delicada.diccionario, 'critico', {})
        tot = getattr(private.informacion_delicada.diccionario, 'total', {})

        lineas = [
            "import os\n",
            "import openpyxl\n\n",
            "test = os.path.join(os.path.dirname(__file__), '..', 'excel base', 'EXPORT.XLSX')\n",
            "if not os.path.exists(test):\n",
            '    test = r"C:\\Users\\Anibal M\\Desktop\\practica 2\\Practica-2\\private\\excel base\\EXPORT.XLSX"\n\n',
            "material = {\n"
        ]
        for k, v in mat.items():
            lineas.append(f"    {json.dumps(str(k))}: {json.dumps(str(v))},\n")
        lineas.append("}\n\n")

        lineas.append("def creardicmar(test):\n")
        lineas.append("    excel = openpyxl.load_workbook(test)\n")
        lineas.append("    hoja = excel.active\n")
        lineas.append("    for i in range(2, hoja.max_row + 1):\n")
        lineas.append('        if hoja.cell(row=i + 1, column=1).value != hoja.cell(row=i, column=1).value and hoja.cell(row = i, column = 2).value != "NULO":\n')
        lineas.append("            codigo = hoja.cell(row=i, column=1).value\n")
        lineas.append("            descripcion = hoja.cell(row=i, column=2).value\n")
        lineas.append("            material[str(codigo)] = str(descripcion)\n")
        lineas.append("    return material\n\n")

        lineas.append("precios = {\n")
        for k, v in pre.items():
            lineas.append(f"    {json.dumps(str(k))}: {json.dumps(str(v))},\n")
        lineas.append("}\n\n")

        lineas.append("ubicaciones = {\n")
        for k, v in ubi.items():
            lineas.append(f"    {json.dumps(str(k))}: {json.dumps(str(v))},\n")
        lineas.append("}\n\n")

        lineas.append("almacenes = {\n")
        for k, v in alm.items():
            lineas.append(f"    {json.dumps(str(k))}: {json.dumps(str(v))},\n")
        lineas.append("}\n\n")

        lineas.append("comprador = {\n")
        for k, v in com.items():
            lineas.append(f"    {json.dumps(str(k))}: {json.dumps(str(v))},\n")
        lineas.append("}\n\n")

        lineas.append("critico = {\n")
        for k, v in cri.items():
            lineas.append(f"    {json.dumps(str(k))}: {json.dumps(str(v))},\n")
        lineas.append("}\n\n")

        lineas.append("total = {\n")
        for k, v in tot.items():
            lineas.append(f"    {json.dumps(str(k))}: {json.dumps(str(v))},\n")
        lineas.append("}\n\n")

        lineas.append("def creardictotal(test):\n")
        lineas.append("    excel = openpyxl.load_workbook(test)\n")
        lineas.append("    hoja = excel.active\n")
        lineas.append("    for i in range(2, hoja.max_row + 1):\n")
        lineas.append('        if str(hoja.cell(row=i, column=3).value).strip() == "M501" and hoja.cell(row=i, column=2).value != "NULO":\n')
        lineas.append("            codigo = hoja.cell(row=i, column=1).value\n")
        lineas.append("            dato = hoja.cell(row=i, column=4).value\n")
        lineas.append("            if codigo is not None:\n")
        lineas.append("                total[str(codigo)] = str(dato) if dato is not None else ''\n")
        lineas.append("    return total\n")

        with open(_RUTA_DICCIONARIO_PY, 'w', encoding='utf-8') as f:
            f.writelines(lineas)
    except Exception as e:
        print(f"Error guardando diccionario.py: {e}")

    # Guardar también copia JSON de respaldo
    for nom in _DIC_VALIDOS:
        jpath = _JSON_FILES.get(nom)
        if jpath:
            try:
                dic_obj = getattr(private.informacion_delicada.diccionario, nom, {})
                with open(jpath, 'w', encoding='utf-8') as f:
                    json.dump(dic_obj, f, ensure_ascii=False, indent=2)
            except Exception:
                pass


def _inicializar_diccionarios():
    """
    Sincroniza los diccionarios vivos si existen archivos JSON guardados previamente.
    Carga total desde Excel si está vacío.
    """
    # Asegurar que existan todos los diccionarios como atributos en el módulo diccionario
    for nom in _DIC_VALIDOS:
        if not hasattr(private.informacion_delicada.diccionario, nom):
            setattr(private.informacion_delicada.diccionario, nom, {})

    modificado = False
    for nombre in _DIC_VALIDOS:
        json_path = _JSON_FILES.get(nombre)
        if json_path and os.path.exists(json_path):
            try:
                with open(json_path, 'r', encoding='utf-8') as f:
                    datos = json.load(f)
                if isinstance(datos, dict) and datos:
                    dic_obj = getattr(private.informacion_delicada.diccionario, nombre, None)
                    if dic_obj is not None:
                        dic_obj.clear()
                        dic_obj.update(datos)
                        modificado = True
            except Exception:
                pass

    # Si total está vacío, cargarlo con la lógica dada desde EXPORT.XLSX si existe
    tot = getattr(private.informacion_delicada.diccionario, 'total', {})
    if not tot:
        test_path = os.path.join(_DIR_INFO, '..', 'excel base', 'EXPORT.XLSX')
        if not os.path.exists(test_path):
            test_path = r"C:\Users\Anibal M\Desktop\practica 2\Practica-2\private\excel base\EXPORT.XLSX"
        if os.path.exists(test_path):
            try:
                if hasattr(private.informacion_delicada.diccionario, 'creardictotal'):
                    private.informacion_delicada.diccionario.creardictotal(test_path)
                else:
                    _cargar_total_desde_excel(tot, test_path)
                modificado = True
            except Exception as e:
                print(f"Error cargando total inicial: {e}")

    if modificado:
        _guardar_en_diccionario_py()

# Ejecutar inicialización al cargar el módulo
_inicializar_diccionarios()


def _validar_nombre(nombre: str) -> str:
    """Valida y devuelve el nombre del diccionario; si es inválido retorna 'ubicaciones'."""
    return nombre if nombre in _DIC_VALIDOS else 'ubicaciones'


def _obtener_diccionario(nombre: str, ruta_excel: str = None) -> dict:
    """
    Obtiene la referencia directa al diccionario en private.informacion_delicada.diccionario.
    """
    nombre = _validar_nombre(nombre)
    if nombre == 'material' and ruta_excel:
        mat = getattr(private.informacion_delicada.diccionario, 'material', {})
        if not mat:
            try:
                private.informacion_delicada.diccionario.creardicmar(ruta_excel)
            except Exception:
                pass
    elif nombre == 'total':
        tot = getattr(private.informacion_delicada.diccionario, 'total', {})
        if not tot:
            archivo = ruta_excel or getattr(private.informacion_delicada.diccionario, 'test', None)
            if not archivo or not os.path.exists(archivo):
                archivo = os.path.join(_DIR_INFO, '..', 'excel base', 'EXPORT.XLSX')
            if archivo and os.path.exists(archivo):
                try:
                    if hasattr(private.informacion_delicada.diccionario, 'creardictotal'):
                        private.informacion_delicada.diccionario.creardictotal(archivo)
                    else:
                        _cargar_total_desde_excel(tot, archivo)
                except Exception as e:
                    print(f"Error cargando total: {e}")
    return getattr(private.informacion_delicada.diccionario, nombre, {})


def diccionarios(ruta):
    """
    Retorna los 3 diccionarios (ubicaciones, precios, materiales) actualizados
    para que cualquier función (como inventario e inventario_preview) use los datos reales.
    """
    dicub = _obtener_diccionario('ubicaciones')
    dicpre = _obtener_diccionario('precios')
    dicmat = _obtener_diccionario('material', ruta)
    return dicub, dicpre, dicmat

# ─── RUTAS PRINCIPALES ───────────────────────────────────────────────────────

@front_flask.route('/')
def home():
    mi_variable = "¡Hola desde Python!"
    ruta = session.get('ruta', '')
    alerta = session.pop('alerta', None)
    return render_template('front.html', dato=mi_variable, dia_hoy=dia, ruta=ruta, alerta=alerta)

@front_flask.route('/ruta', methods=['GET'])
def ruta():
    archivo_seleccionado = pedir_archivo_visual()
    if archivo_seleccionado:
        session['ruta'] = archivo_seleccionado
    return redirect(url_for('home'))

# ─── STOCK TOTAL ─────────────────────────────────────────────────────────────

@front_flask.route('/stock/preview', methods=['GET'])
def stock_preview():
    ruta = session.get('ruta') or request.args.get('ruta')
    if not ruta:
        session['alerta'] = "Advertencia: Primero debes seleccionar un archivo."
        return redirect(url_for('home'))
    columnas, filas = total_preview(ruta)
    return render_template('front.html',
                           dato="¡Hola desde Python!",
                           dia_hoy=dia,
                           ruta=ruta,
                           alerta=None,
                           preview_titulo="Stock Total",
                           preview_columnas=columnas,
                           preview_filas=filas,
                           descarga_url=url_for('stock_descargar'))

@front_flask.route('/stock/descargar', methods=['GET'])
def stock_descargar():
    ruta = session.get('ruta') or request.args.get('ruta')
    if not ruta:
        session['alerta'] = "Advertencia: Primero debes seleccionar un archivo."
        return redirect(url_for('home'))
    nombre_archivo, ruta_creacion = berfre_total(ruta)
    return send_file(ruta_creacion,
                     as_attachment=True,
                     download_name=nombre_archivo,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ─── BODEGA (M501) ───────────────────────────────────────────────────────────

@front_flask.route('/bodega/preview', methods=['GET'])
def bodega_preview():
    ruta = session.get('ruta') or request.args.get('ruta')
    if not ruta:
        session['alerta'] = "Advertencia: Primero debes seleccionar un archivo."
        return redirect(url_for('home'))
    columnas, filas = filtro1_preview(ruta)
    return render_template('front.html',
                           dato="¡Hola desde Python!",
                           dia_hoy=dia,
                           ruta=ruta,
                           alerta=None,
                           preview_titulo="Bodega M501",
                           preview_columnas=columnas,
                           preview_filas=filas,
                           descarga_url=url_for('bodega_descargar'))

@front_flask.route('/bodega/descargar', methods=['GET'])
def bodega_descargar():
    ruta = session.get('ruta') or request.args.get('ruta')
    if not ruta:
        session['alerta'] = "Advertencia: Primero debes seleccionar un archivo."
        return redirect(url_for('home'))
    nombre_archivo, ruta_creacion = filtro1(ruta)
    return send_file(ruta_creacion,
                     as_attachment=True,
                     download_name=nombre_archivo,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ─── INVENTARIO POR LUGAR ─────────────────────────────────────────────────────

@front_flask.route('/inventario_lugar/preview', methods=['GET'])
def inventario_preview_route():
    ruta = session.get('ruta') or request.args.get('ruta')
    if not ruta:
        session['alerta'] = "Advertencia: Primero debes seleccionar un archivo."
        return redirect(url_for('home'))
    dicub, dicpre, dicmat = diccionarios(ruta)
    columnas, filas = inventario_preview(ruta, dicub)
    return render_template('front.html',
                           dato="¡Hola desde Python!",
                           dia_hoy=dia,
                           ruta=ruta,
                           alerta=None,
                           preview_titulo="Inventario por Lugar",
                           preview_columnas=columnas,
                           preview_filas=filas,
                           descarga_url=url_for('inventario_descargar'))

@front_flask.route('/inventario_lugar/descargar', methods=['GET'])
def inventario_descargar():
    ruta = session.get('ruta') or request.args.get('ruta')
    if not ruta:
        session['alerta'] = "Advertencia: Primero debes seleccionar un archivo."
        return redirect(url_for('home'))
    dicub, dicpre, dicmat = diccionarios(ruta)
    nombre_archivo, ruta_creacion = inventario(ruta, dicub)
    return send_file(ruta_creacion,
                     as_attachment=True,
                     download_name=nombre_archivo,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

# ─── DICCIONARIOS (editables por el frontend) ────────────────────────────────

def _dic_importar_excel_flexible(file_bytes: bytes, nombre: str) -> dict:
    """
    Importa entradas desde un archivo Excel en memoria.
    Para 'total', detecta si es formato EXPORT (col 1=código, col 3='M501', col 4=dato).
    En cualquier otro caso, toma col 1=código y col 2=valor.
    """
    buffer = io.BytesIO(file_bytes)
    excel = openpyxl.load_workbook(buffer)
    hoja = excel.active
    resultado = {}

    if nombre == 'total':
        es_export = False
        for r in range(2, min(hoja.max_row + 1, 30)):
            if str(hoja.cell(row=r, column=3).value).strip() == 'M501':
                es_export = True
                break
        if es_export:
            for i in range(2, hoja.max_row + 1):
                c3 = str(hoja.cell(row=i, column=3).value).strip() if hoja.cell(row=i, column=3).value is not None else ""
                c2 = hoja.cell(row=i, column=2).value
                if c3 == 'M501' and c2 != 'NULO':
                    codigo = hoja.cell(row=i, column=1).value
                    dato = hoja.cell(row=i, column=4).value
                    if codigo is not None:
                        resultado[str(codigo)] = str(dato) if dato is not None else ""
            return resultado

    # Formato estándar de 2 columnas (A: Clave, B: Valor)
    for i in range(2, hoja.max_row + 1):
        clave = hoja.cell(row=i, column=1).value
        valor = hoja.cell(row=i, column=2).value
        if clave is not None:
            resultado[str(clave)] = str(valor) if valor is not None else ""
    return resultado


@front_flask.route('/diccionario_mod', methods=['GET'])
def diccionario_mod():
    nombre = _validar_nombre(request.args.get('dic', 'ubicaciones'))
    ruta_excel = session.get('ruta') or request.args.get('ruta')
    dic = _obtener_diccionario(nombre, ruta_excel if nombre in ['material', 'total'] else None)
    alerta = session.pop('alerta', None)
    busqueda = request.args.get('q', '').strip().lower()
    
    # Filtrar sólo para la visualización si hay búsqueda
    if busqueda:
        dic_filtrado = {k: v for k, v in dic.items()
                        if busqueda in str(k).lower() or busqueda in str(v).lower()}
    else:
        dic_filtrado = dic

    return render_template('diccionario_mod.html',
                           dic=dic_filtrado,
                           dic_nombre=nombre,
                           dic_label=_DIC_LABELS[nombre],
                           dic_labels=_DIC_LABELS,
                           alerta=alerta,
                           busqueda=request.args.get('q', ''),
                           ruta_excel=ruta_excel,
                           material_sin_ruta=(nombre == 'material' and not ruta_excel and not dic))


@front_flask.route('/diccionario_mod/agregar', methods=['POST'])
def diccionario_agregar():
    nombre = _validar_nombre(request.form.get('dic_nombre', 'ubicaciones'))
    clave = str(request.form.get('clave', '')).strip()
    valor = str(request.form.get('valor', '')).strip()
    if not clave:
        session['alerta'] = "Error: la clave no puede estar vacía."
        return redirect(url_for('diccionario_mod', dic=nombre))
    
    ruta_excel = session.get('ruta')
    dic = _obtener_diccionario(nombre, ruta_excel if nombre in ['material', 'total'] else None)
    dic[clave] = valor
    _guardar_en_diccionario_py()
    session['alerta'] = f"✔ Entrada '{clave}' agregada correctamente en {nombre}."
    return redirect(url_for('diccionario_mod', dic=nombre))


@front_flask.route('/diccionario_mod/modificar', methods=['POST'])
def diccionario_modificar():
    nombre = _validar_nombre(request.form.get('dic_nombre', 'ubicaciones'))
    clave = str(request.form.get('clave', '')).strip()
    valor = str(request.form.get('valor', '')).strip()
    if not clave:
        session['alerta'] = "Error: la clave no puede estar vacía."
        return redirect(url_for('diccionario_mod', dic=nombre))
    
    ruta_excel = session.get('ruta')
    dic = _obtener_diccionario(nombre, ruta_excel if nombre in ['material', 'total'] else None)
    dic[clave] = valor
    _guardar_en_diccionario_py()
    session['alerta'] = f"✔ Entrada '{clave}' actualizada correctamente."
    return redirect(url_for('diccionario_mod', dic=nombre))


@front_flask.route('/diccionario_mod/eliminar', methods=['POST'])
def diccionario_eliminar():
    nombre = _validar_nombre(request.form.get('dic_nombre', 'ubicaciones'))
    clave = str(request.form.get('clave', '')).strip()
    if not clave:
        session['alerta'] = "Error: la clave no puede estar vacía."
        return redirect(url_for('diccionario_mod', dic=nombre))
    
    ruta_excel = session.get('ruta')
    dic = _obtener_diccionario(nombre, ruta_excel if nombre in ['material', 'total'] else None)
    if clave in dic:
        del dic[clave]
        _guardar_en_diccionario_py()
        session['alerta'] = f"✔ Entrada '{clave}' eliminada correctamente de {nombre}."
    else:
        session['alerta'] = f"⚠ La clave '{clave}' no existe en el diccionario {nombre}."
    return redirect(url_for('diccionario_mod', dic=nombre))


@front_flask.route('/diccionario_mod/exportar', methods=['GET'])
def diccionario_exportar():
    nombre = _validar_nombre(request.args.get('dic', 'ubicaciones'))
    ruta_excel = session.get('ruta')
    dic = _obtener_diccionario(nombre, ruta_excel if nombre in ['material', 'total'] else None)
    buffer = dic_exportar_bytes(dic, nombre)
    return send_file(buffer,
                     as_attachment=True,
                     download_name=f"diccionario_{nombre}_exportado.xlsx",
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@front_flask.route('/diccionario_mod/importar', methods=['POST'])
def diccionario_importar():
    nombre = _validar_nombre(request.form.get('dic_nombre', 'ubicaciones'))
    archivo = request.files.get('archivo')
    if not archivo or archivo.filename == '':
        session['alerta'] = "Error: no se seleccionó ningún archivo."
        return redirect(url_for('diccionario_mod', dic=nombre))
    try:
        nuevas = _dic_importar_excel_flexible(archivo.read(), nombre)
        ruta_excel = session.get('ruta')
        dic = _obtener_diccionario(nombre, ruta_excel if nombre in ['material', 'total'] else None)
        dic.update(nuevas)
        _guardar_en_diccionario_py()
        session['alerta'] = f"✔ Importación exitosa: {len(nuevas)} entradas cargadas en {nombre}."
    except Exception as e:
        session['alerta'] = f"Error al importar: {e}"
    return redirect(url_for('diccionario_mod', dic=nombre))


@front_flask.route('/diccionario_mod/recargar', methods=['POST'])
def diccionario_recargar():
    nombre = _validar_nombre(request.form.get('dic_nombre', 'total'))
    ruta_excel = session.get('ruta') or getattr(private.informacion_delicada.diccionario, 'test', None)
    if not ruta_excel or not os.path.exists(ruta_excel):
        ruta_excel = os.path.join(_DIR_INFO, '..', 'excel base', 'EXPORT.XLSX')

    if not os.path.exists(ruta_excel):
        session['alerta'] = "Error: no se encontró ningún archivo Excel para recargar."
        return redirect(url_for('diccionario_mod', dic=nombre))

    try:
        dic_obj = getattr(private.informacion_delicada.diccionario, nombre, None)
        if nombre == 'total':
            if dic_obj is not None:
                dic_obj.clear()
            if hasattr(private.informacion_delicada.diccionario, 'creardictotal'):
                private.informacion_delicada.diccionario.creardictotal(ruta_excel)
            else:
                _cargar_total_desde_excel(dic_obj, ruta_excel)
            _guardar_en_diccionario_py()
            session['alerta'] = f"✔ Diccionario Total recargado exitosamente ({len(dic_obj)} entradas cargadas)."
        elif nombre == 'material':
            if hasattr(private.informacion_delicada.diccionario, 'creardicmar'):
                private.informacion_delicada.diccionario.creardicmar(ruta_excel)
            _guardar_en_diccionario_py()
            session['alerta'] = f"✔ Diccionario Materiales recargado exitosamente ({len(dic_obj)} entradas cargadas)."
        else:
            session['alerta'] = f"La recarga automática solo aplica a 'total' o 'material'."
    except Exception as e:
        session['alerta'] = f"Error al recargar {nombre}: {e}"

    return redirect(url_for('diccionario_mod', dic=nombre))

# ─────────────────────────────────────────────────────────────────────────────

if __name__ == '__main__':
    front_flask.run(debug=True, port=5000)